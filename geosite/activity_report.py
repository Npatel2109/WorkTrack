import io
import re
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.datetime import to_excel

from .models import REPORT_COMPANY_NAME
from .timeutils import india_date

ACTIVITY_REPORT_TEMPLATE = (
    Path(__file__).parent / "excel_templates" / "activity_report.xlsx"
)


def _fill_template(rows, site, employee):
    # Patch only worksheet values so Excel-specific parts survive unchanged.
    template_bytes = ACTIVITY_REPORT_TEMPLATE.read_bytes()
    workbook = load_workbook(io.BytesIO(template_bytes))
    sheet = workbook["SAR-R0"]
    if workbook.sheetnames[0] != "SAR-R0":
        raise ValueError("Activity template must have SAR-R0 as its first worksheet.")
    remarks_row = next(
        (cell.row for cell in sheet["A"] if cell.value == "REMARKS"), None
    )
    if remarks_row is None or remarks_row <= 9:
        raise ValueError("Activity template must have activity rows from row 9 followed by REMARKS.")
    extra = max(0, len(rows) - (remarks_row - 9))
    values = {}

    def text(address, value):
        # Use openpyxl's string validation, but never treat user data as formulas.
        value = Cell(sheet).check_string(str(value))
        values[address] = (
            '<is><t xml:space="preserve">'
            + escape(value)
            + "</t></is>"
        )

    def number(address, value):
        values[address] = f"<v>{value}</v>"

    def fill_label(address, value, destination=None):
        label = sheet[address].value
        if not isinstance(label, str) or ":" not in label:
            raise ValueError(f"Activity template is missing a label at {address}.")
        prefix, _, suffix = label.partition(":")
        trailing = suffix[len(suffix.rstrip()):].strip(" \t")
        text(destination or address, f"{prefix}: {value}{trailing}")

    fill_label("A3", site.project.name if site.project else site.project_name or "Not specified")
    fill_label("A4", site.customer_name)
    fill_label("A5", f"{datetime.now(timezone.utc):%d %b %Y}")
    fill_label("A6", site.site_name)
    fill_label("A7", employee.name)
    for index, (day, days, detail) in enumerate(rows, 9):
        number(f"A{index}", index - 8)
        text(f"B{index}", site.site_name)
        number(f"C{index}", to_excel(day, workbook.epoch))
        number(f"D{index}", days)
        text(f"E{index}", detail)
    remarks = site.report_remarks.splitlines()
    for offset in range(1, 6):
        text(f"C{remarks_row + extra + offset}", remarks[offset - 1] if offset <= len(remarks) else "")
    number(f"K{remarks_row + extra + 6}", sum(days for _, days, _ in rows))
    # Read labels at their original positions before shifting the footer XML.
    footer_values = (
        (f"A{remarks_row + 7}", site.customer_name),
        (f"I{remarks_row + 7}", REPORT_COMPANY_NAME),
        (f"A{remarks_row + 8}", site.commissioning_manager or ""),
        (f"E{remarks_row + 8}", site.project_manager or ""),
        (f"I{remarks_row + 8}", employee.name),
    )
    for address, value in footer_values:
        match = re.fullmatch(r"([A-Z]+)(\d+)", address)
        destination = f"{match[1]}{int(match[2]) + extra}"
        if address.startswith("I") and address == f"I{remarks_row + 7}":
            text(destination, f"M/S {value}")
        else:
            fill_label(address, value, destination)
    workbook.close()

    def shift_reference(match):
        row = int(match[2])
        return f"{match[1]}{row + extra if row >= remarks_row else row}"

    output = io.BytesIO()
    with ZipFile(io.BytesIO(template_bytes)) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                if extra:
                    row_pattern = r'<row\b[^>]*\br="(\d+)"[^>]*>.*?</row>'
                    activity_row = next(
                        match[0] for match in re.finditer(row_pattern, xml, re.DOTALL)
                        if int(match[1]) == remarks_row - 1
                    )

                    def shift_row(match):
                        if int(match[1]) < remarks_row:
                            return match[0]
                        shifted = re.sub(r"(\$?[A-Z]+\$?)(\d+)", shift_reference, match[0])
                        return re.sub(r'(<row\b[^>]*\br=")\d+', lambda m: m[1] + str(int(match[1]) + extra), shifted, count=1)

                    xml = re.sub(row_pattern, shift_row, xml, flags=re.DOTALL)
                    added_rows = []
                    for row in range(remarks_row, remarks_row + extra):
                        clone = re.sub(r'(<row\b[^>]*\br=")\d+', lambda m: m[1] + str(row), activity_row, count=1)
                        clone = re.sub(r'(\br="[A-Z]+)\d+(")', lambda m: m[1] + str(row) + m[2], clone)
                        added_rows.append(clone)
                    xml = re.sub(
                        rf'<row\b(?=[^>]*\br="{remarks_row + extra}")',
                        lambda m: "".join(added_rows) + m[0], xml, count=1,
                    )
                    xml = re.sub(
                        r'(<(?:mergeCell|dimension)\b[^>]*\bref=")([^"]+)(")',
                        lambda m: m[1] + re.sub(r"(\$?[A-Z]+\$?)(\d+)", shift_reference, m[2]) + m[3],
                        xml,
                    )
                    merges = "".join(f'<mergeCell ref="E{row}:L{row}"/>' for row in range(remarks_row, remarks_row + extra))
                    xml = xml.replace("</mergeCells>", merges + "</mergeCells>")
                    xml = re.sub(r'(<mergeCells count=")(\d+)', lambda m: m[1] + str(int(m[2]) + extra), xml)

                def fill_cell(match):
                    address = match[1]
                    if address not in values:
                        return match[0]
                    opening = match[0].split(">", 1)[0].rstrip("/")
                    opening = re.sub(r'\s+t="[^"]*"', "", opening)
                    content = values.pop(address)
                    cell_type = ' t="inlineStr"' if content.startswith("<is>") else ""
                    return opening + cell_type + ">" + content + "</c>"

                xml = re.sub(
                    r'<c\b(?=[^>]*\br="([A-Z]+\d+)")[^>]*?(?:/>|>.*?</c>)',
                    fill_cell, xml, flags=re.DOTALL,
                )
                if values:
                    raise ValueError(f"Activity template is missing data cells: {', '.join(values)}")
                data = xml.encode("utf-8")
            elif entry.filename == "xl/workbook.xml" and extra:
                xml = data.decode("utf-8")
                xml = re.sub(
                    r'(<definedName\b[^>]*name="_xlnm.Print_Area"[^>]*>)(.*?)(</definedName>)',
                    lambda m: m[1] + re.sub(r"(\$?[A-Z]+\$?)(\d+)", shift_reference, m[2]) + m[3],
                    xml,
                )
                data = xml.encode("utf-8")
            target.writestr(entry, data)
    output.seek(0)
    return output


def activity_workbook(site, visits, leaves, start, end, employee, automatic_leave=()):
    by_date = defaultdict(list)
    for visit in visits:
        by_date[india_date(visit.checkout_time)].append(visit)
    rows = []
    day = start
    while day <= end:
        billable_day_recorded = False
        for visit in by_date[day]:
            report = visit.report
            detail = report.work_performed
            if report.is_billable is None:
                legacy = [
                    report.activities_completed, report.issues_found,
                    report.recommendations, report.materials_used, report.remarks,
                ]
                detail = "\n".join([detail] + [value for value in legacy if value])
                detail += "\n[Billing not recorded in legacy report; counted as 0]"
            elif not report.is_billable:
                detail += "\n[Non-billable]"
            billable_day = report.is_billable is True and not billable_day_recorded
            billable_day_recorded = billable_day_recorded or billable_day
            rows.append((day, int(billable_day), detail))
        leave_rows = [leave for leave in leaves if leave.start_date <= day <= leave.end_date]
        for leave in leave_rows:
            reason = leave.reason
            if leave.period != "full":
                reason = f"{leave.period_label}: {reason}"
            rows.append((day, 0, reason))
        automatic_rows = [item for item in automatic_leave if item["date"] == day]
        for item in automatic_rows:
            rows.append((day, 0, f"Automatic leave ({item['days']:g} day): no site visit or approved work confirmation."))
        if not by_date[day] and not leave_rows and not automatic_rows:
            rows.append((day, 0, "Sunday" if day.weekday() == 6 else "No activity recorded"))
        day += timedelta(days=1)
    return _fill_template(rows, site, employee)
