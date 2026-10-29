import base64
import io
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from unittest.mock import patch
from urllib.parse import urlparse
from datetime import date, datetime, timedelta, timezone
from html.parser import HTMLParser

from PIL import Image
from pypdf import PdfReader
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from sqlalchemy.exc import IntegrityError

from geosite import create_app
from geosite.attendance import as_utc, close_overdue_visits, daily_cutoff
from geosite.timeutils import INDIA_TIME, india_date
from geosite.extensions import db
from geosite.models import AssignedSite, Attendance, AuditLog, Invoice, LeaveRequest, Project, REPORT_COMPANY_NAME, SiteAssignment, SiteReport, User


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.uploads = tempfile.TemporaryDirectory()
        self.app = create_app(
            {
                "TESTING": True,
                "SECRET_KEY": "test-secret",
                "SQLALCHEMY_DATABASE_URI": "sqlite://",
                "SQLALCHEMY_TRACK_MODIFICATIONS": False,
                "UPLOAD_FOLDER": self.uploads.name,
                "WTF_CSRF_ENABLED": False,
            }
        )
        self.context = self.app.app_context()
        self.context.push()
        db.create_all()
        self.admin = User(
            name="Admin",
            email="admin@example.com",
            mobile="100",
            designation="Administrator",
            role="admin",
        )
        self.admin.set_password("admin-password-123")
        self.employee = User(
            name="Engineer",
            email="engineer@example.com",
            mobile="200",
            designation="Field Engineer",
        )
        self.employee.set_password("engineer-password-123")
        self.project = Project(project_number="PRJ-002", name="Project Two")
        db.session.add_all([self.admin, self.employee, self.project])
        db.session.flush()
        self.site = AssignedSite(
            site_name="North Site",
            customer_name="Example Customer",
            address="1 Example Road",
            latitude=0,
            longitude=0,
            allowed_radius=100,
        )
        db.session.add(self.site)
        db.session.flush()
        db.session.add(SiteAssignment(employee=self.employee, site=self.site))
        db.session.commit()
        self.client = self.app.test_client()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        db.engine.dispose()
        self.context.pop()
        self.uploads.cleanup()

    @staticmethod
    def _image_data():
        buffer = io.BytesIO()
        Image.new("RGB", (20, 20), "white").save(buffer, format="PNG")
        return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode()

    def _login_employee(self):
        return self.client.post(
            "/login",
            data={"email": self.employee.email, "password": "engineer-password-123"},
            follow_redirects=True,
        )

    def test_employee_cannot_open_admin_page(self):
        self._login_employee()
        self.assertEqual(self.client.get("/admin").status_code, 403)

    def test_activity_excel_layout_billing_leave_and_missing_dates(self):
        start = date(2026, 10, 19)
        self.site.project_name = "=Project text"
        self.site.report_company = REPORT_COMPANY_NAME
        self.site.commissioning_manager = "Commissioning Lead"
        self.site.project_manager = "Project Lead"
        self.site.report_remarks = "First remark\n=literal remark"
        for billing, detail in ((True, "Billable work"), (False, "Non-billable meeting"),
                                (True, "=literal activity")):
            visit = Attendance(
                employee_id=self.employee.id, site_id=self.site.id,
                status="completed", checkin_lat=0, checkin_lon=0,
                checkin_selfie="selfie.png",
                checkout_time=datetime(2026, 10, 19, 12, tzinfo=timezone.utc),
            )
            db.session.add(visit)
            db.session.flush()
            db.session.add(SiteReport(
                attendance=visit, work_performed=detail, is_billable=billing,
                activities_completed="", issues_found="", recommendations="",
                materials_used="", remarks="", engineer_name="Engineer",
                customer_representative_name="Customer", employee_signature="a.png",
                customer_signature="b.png",
            ))
        db.session.add(LeaveRequest(
            employee_id=self.employee.id, start_date=start + timedelta(days=1),
            end_date=start + timedelta(days=1), period="full", reason="Sick Leave",
            status="approved",
        ))
        db.session.add(LeaveRequest(
            employee_id=self.employee.id, start_date=start + timedelta(days=2),
            end_date=start + timedelta(days=2), period="full", reason="Unapproved reason",
            status="pending",
        ))
        db.session.commit()
        self._login_employee()
        response = self.client.get("/reports/activity.xlsx", query_string={
            "site_id": self.site.id, "start_date": start.isoformat(),
            "end_date": (start + timedelta(days=6)).isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.data)).active
        self.assertEqual(sheet.title, "SAR-R0")
        self.assertEqual(sheet["A1"].value, "SITE ACTIVITY REPORT")
        self.assertEqual(sheet["A3"].value, "PROJECT : =Project text")
        self.assertTrue(sheet["A5"].value.startswith("REPORT UPDATE DATE :"))
        self.assertFalse(any("PO NO" in str(cell.value) or "PO DATE" in str(cell.value)
                             for row in sheet for cell in row))
        self.assertIn("E8:L8", str(sheet.merged_cells))
        self.assertEqual([sheet.cell(8, c).value for c in range(1, 6)],
                         ["SR NO", "SITE", "DATE", "DAY", "ACTIVITY DETAIL"])
        self.assertEqual([sheet.cell(r, 4).value for r in range(9, 13)], [1, 0, 0, 0])
        self.assertEqual(sheet["E9"].value, "Billable work\nNon-billable meeting\n[Non-billable]\n=literal activity")
        self.assertEqual(sheet["E9"].data_type, "s")
        self.assertTrue(sheet["E9"].alignment.wrap_text)
        self.assertGreater(sheet.row_dimensions[9].height, 37.95)
        self.assertEqual(sheet["E10"].value, "Sick Leave")
        self.assertEqual(sheet["E11"].value, "No activity recorded")
        self.assertEqual(sheet["E15"].value, "Sunday")
        self.assertEqual(sheet["K31"].value, 1)
        self.assertEqual(sheet["A25"].value, "REMARKS")
        self.assertEqual([sheet.cell(r, 1).value for r in range(26, 31)], [1, 2, 3, 4, 5])
        self.assertEqual(sheet["C26"].value, "First remark")
        self.assertEqual(sheet["C27"].value, "=literal remark")
        self.assertEqual(sheet["C27"].data_type, "s")
        self.assertIn(sheet["C28"].value, (None, ""))
        self.assertEqual(sheet["I32"].value, f"M/S {REPORT_COMPANY_NAME}")
        self.assertIn("Commissioning Manager : Commissioning Lead", sheet["A33"].value)
        self.assertIn("Project Manager : Project Lead", sheet["E33"].value)
        self.assertIn("Engineer : Engineer", sheet["I33"].value)
        for merge in ("A33:D37", "E33:H37", "I33:L37", "C26:L26"):
            self.assertIn(merge, str(sheet.merged_cells))
        self.assertEqual(sheet.page_setup.orientation, "portrait")
        self.assertNotIn("Unapproved reason", [cell.value for row in sheet for cell in row])
        self.assertEqual(sheet.freeze_panes, "E9")

    def test_activity_excel_uses_replaceable_template_formatting(self):
        from geosite.activity_report import ACTIVITY_REPORT_TEMPLATE

        template_path = Path(self.uploads.name) / "custom_activity_report.xlsx"
        template = load_workbook(ACTIVITY_REPORT_TEMPLATE)
        template.active["A8"].fill = PatternFill(
            fill_type="solid", fgColor="FFFF0000"
        )
        template.active.column_dimensions["A"].width = 32
        template.save(template_path)
        template.close()

        self._login_employee()
        with patch(
            "geosite.activity_report.ACTIVITY_REPORT_TEMPLATE", template_path
        ):
            response = self.client.get("/reports/activity.xlsx", query_string={
                "site_id": self.site.id,
                "start_date": "2026-10-19",
                "end_date": "2026-10-19",
            })

        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.data)).active
        self.assertEqual(sheet["A8"].fill.fgColor.rgb, "FFFF0000")
        self.assertEqual(sheet.column_dimensions["A"].width, 32)

    def test_activity_excel_preserves_template_and_expands_only_when_needed(self):
        from geosite.activity_report import ACTIVITY_REPORT_TEMPLATE

        self._login_employee()
        for days in (1, 16, 20):
            with self.subTest(days=days):
                response = self.client.get("/reports/activity.xlsx", query_string={
                    "site_id": self.site.id,
                    "start_date": "2026-10-01",
                    "end_date": f"2026-10-{days:02}",
                })
                self.assertEqual(response.status_code, 200)
                sheet = load_workbook(io.BytesIO(response.data)).active
                extra = max(0, days - 16)
                self.assertEqual(sheet.max_row, 37 + extra)
                self.assertEqual(sheet[f"A{25 + extra}"].value, "REMARKS")
                self.assertEqual(sheet[f"A{8 + days}"].value, days)
                self.assertEqual(sheet[f"I{33 + extra}"].value.strip(), "Engineer : Engineer")
                self.assertIn(f"E{8 + days}:L{8 + days}", str(sheet.merged_cells))
                with ZipFile(ACTIVITY_REPORT_TEMPLATE) as template, ZipFile(io.BytesIO(response.data)) as exported:
                    self.assertEqual(template.namelist(), exported.namelist())
                    changed = {"xl/worksheets/sheet1.xml"}
                    if extra:
                        changed.add("xl/workbook.xml")
                    for part in template.namelist():
                        if part not in changed:
                            self.assertEqual(template.read(part), exported.read(part), part)
                original = load_workbook(ACTIVITY_REPORT_TEMPLATE).active
                if not extra:
                    self.assertEqual(str(sheet.merged_cells), str(original.merged_cells))
                    self.assertEqual(str(sheet.print_area), str(original.print_area))
                    for row in original:
                        for cell in row:
                            self.assertEqual(sheet[cell.coordinate].style_id, cell.style_id)
                        self.assertEqual(sheet.row_dimensions[row[0].row].height,
                                         original.row_dimensions[row[0].row].height)

    def test_activity_excel_enforces_owner_and_valid_date_range(self):
        other = User(name="Other", email="other@example.com", mobile="3")
        other.set_password("other-password-123")
        db.session.add(other)
        db.session.flush()
        site = AssignedSite(site_name="Private Site",
                            customer_name="Customer", address="Address",
                            latitude=0, longitude=0)
        db.session.add(site)
        db.session.flush()
        db.session.add(SiteAssignment(employee=other, site=site))
        db.session.commit()
        self._login_employee()
        params = {"site_id": site.id, "start_date": "2026-10-19", "end_date": "2026-10-20"}
        self.assertEqual(self.client.get("/reports/activity.xlsx", query_string=params).status_code, 404)
        self.assertNotIn(b"Private Site", self.client.get("/reports").data)
        for updates in ({"start_date": "bad"}, {"end_date": "2026-10-01"},
                        {"end_date": "2028-10-20"}, {"site_id": "bad"}):
            response = self.client.get("/reports/activity.xlsx", query_string={**params, **updates},
                                       follow_redirects=True)
            self.assertIn(b"Select a site and valid", response.data)
        self.client.post("/logout")
        self.client.post("/login", data={"email": self.admin.email, "password": "admin-password-123"})
        params["employee_id"] = other.id
        self.assertEqual(self.client.get("/reports/activity.xlsx", query_string=params).status_code, 200)

    def test_activity_excel_counts_multiple_visits_on_one_india_date_once(self):
        self._completed_display_visit(datetime(2026, 10, 19, 20, tzinfo=timezone.utc))
        self._completed_display_visit(datetime(2026, 10, 20, 3, tzinfo=timezone.utc))
        self._login_employee()

        response = self.client.get("/reports/activity.xlsx", query_string={
            "site_id": self.site.id,
            "start_date": "2026-10-20",
            "end_date": "2026-10-20",
        })

        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.data)).active
        self.assertEqual(sheet["C9"].value.date().isoformat(), "2026-10-20")
        self.assertIsNone(sheet["C10"].value)
        self.assertEqual(sheet["D9"].value, 1)
        self.assertEqual(sheet["E9"].value.count("\n"), 1)
        total_row = next(row[0].row for row in sheet.iter_rows() if row[0].value == "TOTAL DAYS :")
        self.assertEqual(sheet.cell(total_row, 11).value, 1)

    def test_admin_and_employee_excel_use_identical_layout_and_data(self):
        params = {"site_id": self.site.id, "start_date": "2026-10-19", "end_date": "2026-10-21"}
        self._login_employee()
        employee_response = self.client.get("/reports/activity.xlsx", query_string=params)
        employee_sheet = load_workbook(io.BytesIO(employee_response.data)).active
        self.assertNotIn(b'id="activity-employee"', self.client.get("/reports").data)
        self.client.post("/logout")
        self.client.post("/login", data={"email": self.admin.email, "password": "admin-password-123"})
        self.assertIn(b'id="activity-employee"', self.client.get("/reports").data)
        response = self.client.get("/admin/reports/export.xlsx", query_string=params, follow_redirects=True)
        self.assertIn(b"Select an employee", response.data)
        params["employee_id"] = self.employee.id
        for route in ("/reports/activity.xlsx", "/admin/reports/export.xlsx"):
            response = self.client.get(route, query_string=params)
            self.assertEqual(response.status_code, 200)
            admin_sheet = load_workbook(io.BytesIO(response.data)).active
            self.assertEqual(list(admin_sheet.values), list(employee_sheet.values))
            self.assertEqual(str(admin_sheet.merged_cells), str(employee_sheet.merged_cells))
            for row in employee_sheet:
                for cell in row:
                    self.assertEqual(admin_sheet[cell.coordinate].style_id, cell.style_id)
        params["employee_id"] = self.admin.id
        self.assertEqual(self.client.get("/reports/activity.xlsx", query_string=params).status_code, 404)

    def test_admin_project_header_changes(self):
        self._login_employee()
        url = f"/admin/sites/{self.site.id}/report-details"
        self.assertEqual(self.client.post(url, data={}).status_code, 403)
        self.client.post("/logout")
        self.client.post("/login", data={"email": self.admin.email, "password": "admin-password-123"})
        project = Project(project_number="PRJ-003", name="Commissioning")
        db.session.add(project)
        db.session.commit()
        self.client.post(url, data={"project_id": str(project.id), "report_company": "Company",
                                   "commissioning_manager": "Manager One",
                                   "project_manager": "Manager Two", "report_remarks": "One\nTwo"})
        db.session.refresh(self.site)
        self.assertEqual(self.site.project_name, "Commissioning")
        self.assertEqual(self.site.project_id, project.id)
        self.assertEqual(self.site.report_company, REPORT_COMPANY_NAME)
        self.assertEqual(self.site.commissioning_manager, "Manager One")
        self.assertEqual(self.site.project_manager, "Manager Two")
        self.assertEqual(self.site.report_remarks, "One\nTwo")
        for values, expected_message in (
            ({"report_remarks": "Line\n" * 6}, b"five remarks lines"),
            ({"project_manager": "x" * 121}, b"120 characters"),
            ({"commissioning_manager": "x" * 121}, b"120 characters"),
        ):
            response = self.client.post(url, data=values, follow_redirects=True)
            self.assertIn(expected_message, response.data)
            db.session.refresh(self.site)
            self.assertEqual(self.site.report_company, REPORT_COMPANY_NAME)
        response = self.client.post(url, data={"project_id": "not-an-id"}, follow_redirects=True)
        self.assertIn(b"Select a valid project", response.data)
        db.session.refresh(self.site)
        self.assertEqual(self.site.project_id, project.id)

    def test_end_to_end_visit_and_merged_report_pdf(self):
        self._login_employee()
        response = self.client.post(f"/visit/start/{self.site.id}")
        self.assertEqual(response.status_code, 302)
        too_far = self.client.post(
            f"/api/visit/{self.site.id}/verify-checkin",
            json={"latitude": 0, "longitude": 0.01},
        )
        self.assertTrue(too_far.get_json()["verified"])
        self.assertFalse(too_far.get_json()["within_radius"])
        self.assertGreater(too_far.get_json()["distance"], 1100)
        self.assertEqual(self.client.get(f"/visit/{self.site.id}/camera").status_code, 200)
        verified = self.client.post(
            f"/api/visit/{self.site.id}/verify-checkin",
            json={"latitude": 0, "longitude": 0},
        )
        self.assertTrue(verified.get_json()["verified"])
        result = self.client.post(
            f"/visit/{self.site.id}/checkin",
            data={"selfie": self._image_data()},
            follow_redirects=True,
        )
        self.assertEqual(result.status_code, 200)
        attendance = Attendance.query.one()
        self.assertEqual(attendance.status, "onsite")
        logout_response = self.client.post("/logout")
        self.assertTrue(logout_response.location.endswith("/login"))
        self.assertEqual(Attendance.query.one().status, "onsite")
        self._login_employee()

        checkout = self.client.post(
            f"/api/visit/{attendance.id}/verify-checkout",
            json={"latitude": 0, "longitude": 0},
        )
        self.assertTrue(checkout.get_json()["verified"])
        payload = {
            "activity_detail": "Inspection and functional checks",
            "billing": "billable",
            "work_performed": "Inspection",
            "activities_completed": "Functional checks",
            "issues_found": "None",
            "recommendations": "Continue monthly checks",
            "materials_used": "None",
            "remarks": "All checks passed",
            "engineer_name": "Engineer",
        }
        for billing in ("", "invalid"):
            response = self.client.post(
                f"/visit/{attendance.id}/report", data={**payload, "billing": billing}
            )
            self.assertIn(b"select Billable or Non-billable", response.data)
            self.assertEqual(SiteReport.query.count(), 0)
        response = self.client.post(
            f"/visit/{attendance.id}/report", data={**payload, "activity_detail": ""}
        )
        self.assertIn(b"Enter activity detail", response.data)
        for name in ("work_performed", "activities_completed", "issues_found",
                     "recommendations", "materials_used", "remarks"):
            payload.pop(name)
        form = self.client.get(f"/visit/{attendance.id}/report")
        self.assertIn(b'id="report-dictate" type="button"', form.data)
        self.assertIn(b"English voice typing", form.data)
        self.assertIn(b"js/report-dictation.js", form.data)
        self.assertIn(b'id="dictation-status" role="status" aria-live="polite"', form.data)
        self.assertNotIn(b'name="customer_representative_name"', form.data)
        self.assertNotIn(b'class="signature-pad"', form.data)
        self.assertNotIn(b'name="employee_signature"', form.data)
        self.assertNotIn(b'name="customer_signature"', form.data)
        response = self.client.post(
            f"/visit/{attendance.id}/report", data=payload, follow_redirects=True
        )
        self.assertEqual(response.status_code, 200)
        db.session.refresh(attendance)
        self.assertEqual(attendance.status, "completed")
        self.assertIsNotNone(SiteReport.query.one())
        self.assertTrue(SiteReport.query.one().is_billable)
        self.assertEqual(SiteReport.query.one().work_performed, "Inspection and functional checks")
        self.assertEqual(SiteReport.query.one().customer_representative_name, "")
        self.assertEqual(SiteReport.query.one().employee_signature, "")
        self.assertEqual(SiteReport.query.one().customer_signature, "")

        self.assertEqual(self.client.get(f"/reports/{attendance.id}/pdf").status_code, 403)
        self.assertEqual(self.client.get("/reports/combined.pdf").status_code, 403)
        self.assertNotIn(b"Download PDF", self.client.get("/reports").data)
        self._login_admin()
        single = self.client.get(f"/reports/{attendance.id}/pdf")
        pdf = PdfReader(io.BytesIO(single.data))
        page_count = len(pdf.pages)
        pdf_text = "\n".join(page.extract_text() for page in pdf.pages)
        for label in ("Customer Representative", "Engineer Signature", "Customer Signature",
                      "Offline Approval", "Date:"):
            self.assertIn(label, pdf_text)
        single.close()
        combined = self.client.get(
            "/reports/combined.pdf",
            query_string={
                "start_date": india_date(attendance.checkout_time).isoformat(),
                "end_date": india_date(attendance.checkout_time).isoformat(),
            },
        )
        self.assertEqual(combined.status_code, 200)
        self.assertEqual(len(PdfReader(io.BytesIO(combined.data)).pages), page_count)

    def test_checkout_requires_activity_billing_and_engineer_name(self):
        self._login_employee()
        self.client.post(
            f"/api/visit/{self.site.id}/verify-checkin",
            json={"latitude": 0, "longitude": 0},
        )
        self.client.post(
            f"/visit/{self.site.id}/checkin", data={"selfie": self._image_data()}
        )
        attendance = Attendance.query.one()
        self.client.post(
            f"/api/visit/{attendance.id}/verify-checkout",
            json={"latitude": 0, "longitude": 0},
        )
        self.client.post(f"/visit/{attendance.id}/report", data={"work_performed": "Inspection"})
        self.assertEqual(Attendance.query.one().status, "report_pending")
        self.assertEqual(SiteReport.query.count(), 0)
        response = self.client.post(f"/visit/{attendance.id}/report", data={
            "activity_detail": "Inspection", "billing": "non_billable", "engineer_name": "",
        })
        self.assertIn(b"enter the engineer name", response.data)
        self.assertEqual(SiteReport.query.count(), 0)
        self.client.post(f"/visit/{attendance.id}/report", data={
            "activity_detail": "Inspection", "billing": "non_billable", "engineer_name": "Engineer",
        })
        self.assertEqual(Attendance.query.one().status, "completed")
        self.assertFalse(SiteReport.query.one().is_billable)

    def _start_location_visit(self, longitude=0.01):
        self._login_employee()
        response = self.client.post(
            f"/api/visit/{self.site.id}/verify-checkin",
            json={"latitude": 0, "longitude": longitude},
        )
        self.assertEqual(response.status_code, 200)
        self.client.post(
            f"/visit/{self.site.id}/checkin", data={"selfie": self._image_data()}
        )
        return Attendance.query.one()

    def test_employee_can_sign_out_with_active_site_visit(self):
        visit = self._start_location_visit()

        response = self.client.post("/logout")

        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.location.endswith("/login"))
        self.assertEqual(self.client.get("/employee").status_code, 302)
        db.session.refresh(visit)
        self.assertEqual(visit.status, "onsite")
        self.assertEqual(AuditLog.query.filter_by(action="logout").count(), 1)

    def test_outside_radius_start_and_stop_record_distances_and_fresh_location(self):
        visit = self._start_location_visit()
        self.assertAlmostEqual(visit.checkin_distance_m, 1111.949, places=2)
        self.assertEqual((visit.site_latitude, visit.site_longitude), (0, 0))
        self.assertEqual(visit.checkin_lon, 0.01)
        self.assertIsNotNone(visit.checkin_location_time)
        response = self.client.post(
            f"/api/visit/{visit.id}/verify-checkout",
            json={"latitude": 0, "longitude": 0.02},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["location_source"], "current")
        db.session.refresh(visit)
        self.assertEqual(visit.status, "report_pending")
        self.assertEqual(visit.checkout_lon, 0.02)
        self.assertAlmostEqual(visit.checkout_distance_m, 2223.899, places=2)
        self.assertEqual(visit.checkout_location_time, visit.checkout_time)
        self.assertFalse(visit.auto_closed)
        self.assertIn(b"Complete report", self.client.get("/employee").data)
        self.assertEqual(self.client.get(f"/reports/{visit.id}/pdf").status_code, 403)
        # Assignment edits must not rewrite the original distance reference.
        self.site.longitude = 1
        db.session.commit()
        self.assertEqual(visit.site_longitude, 0)
        self.assertAlmostEqual(visit.checkin_distance_m, 1111.949, places=2)

    def test_location_off_stop_uses_latest_persisted_fix_not_client_coordinates(self):
        visit = self._start_location_visit()
        response = self.client.post(
            f"/api/visit/{visit.id}/location", json={"latitude": 0, "longitude": 0.03}
        )
        self.assertEqual(response.status_code, 200)
        db.session.refresh(visit)
        captured = visit.last_location_time
        response = self.client.post(f"/api/visit/{visit.id}/verify-checkout", json={
            "location_unavailable": True, "reason": "permission_denied",
            "latitude": 50, "longitude": 50,
        })
        self.assertEqual(response.status_code, 200)
        db.session.refresh(visit)
        self.assertEqual(visit.checkout_lon, 0.03)
        self.assertEqual(visit.checkout_location_time, captured)
        self.assertEqual(visit.checkout_location_source, "last_known")
        self.assertEqual(visit.checkout_location_reason, "permission_denied")
        self.assertAlmostEqual(visit.checkout_distance_m, 3335.848, places=2)
        self.assertEqual(AuditLog.query.filter_by(action="site_checkout").count(), 1)
        stopped_at = visit.checkout_time
        self.client.post(f"/api/visit/{visit.id}/verify-checkout", json={"latitude": 1, "longitude": 1})
        db.session.refresh(visit)
        self.assertEqual(visit.checkout_time, stopped_at)
        self.assertEqual(visit.checkout_lon, 0.03)
        self.assertEqual(AuditLog.query.filter_by(action="site_checkout").count(), 1)
        response = self.client.post(f"/api/visit/{visit.id}/location", json={"latitude": 1, "longitude": 1})
        self.assertEqual(response.status_code, 409)
        db.session.refresh(visit)
        self.assertEqual(visit.last_location_lon, 0.03)

    def test_gps_failures_fall_back_to_start_and_invalid_input_is_explicit(self):
        visit = self._start_location_visit()
        for payload in ([], ["bad"], {"latitude": 91, "longitude": 0},
                        {"latitude": 0, "longitude": float("nan")},
                        {"location_unavailable": True, "reason": "bad"}):
            response = self.client.post(f"/api/visit/{visit.id}/verify-checkout", json=payload)
            self.assertEqual(response.status_code, 400)
            self.assertIn("error", response.json)
            self.assertEqual(visit.status, "onsite")
        response = self.client.post(f"/api/visit/{visit.id}/verify-checkout", json={
            "location_unavailable": True, "reason": "timeout",
        })
        self.assertEqual(response.status_code, 200)
        db.session.refresh(visit)
        self.assertEqual(visit.checkout_lon, visit.checkin_lon)
        self.assertEqual(visit.checkout_location_source, "last_known")

    def test_half_day_leave_allows_selected_site_gps_and_selfie_start(self):
        today = datetime.now(INDIA_TIME).date()
        db.session.add(LeaveRequest(
            employee_id=self.employee.id, start_date=today, end_date=today,
            period="morning", reason="Approved half day", status="approved",
        ))
        db.session.commit()
        self._login_employee()
        response = self.client.post("/employee/work-status",
                                    data={"work_type": "onsite", "site_id": self.site.id})
        self.assertEqual(response.location, f"/visit/{self.site.id}/checkin-location")
        self.assertEqual(self.client.post(f"/api/visit/{self.site.id}/verify-checkin",
                                         json={"latitude": 0, "longitude": 0}).status_code, 200)
        self.assertEqual(self.client.get(f"/visit/{self.site.id}/camera").status_code, 200)
        self.client.post(f"/visit/{self.site.id}/checkin", data={"selfie": self._image_data()})
        visit = Attendance.query.one()
        self.assertEqual(visit.status, "onsite")
        self.assertEqual(visit.site_id, self.site.id)
        self.assertEqual(visit.employee_id, self.employee.id)

    def test_start_requires_valid_fresh_location_and_live_selfie(self):
        self._login_employee()
        self.assertEqual(self.client.get(f"/visit/{self.site.id}/camera").status_code, 403)
        self.assertEqual(self.client.post(
            f"/api/visit/{self.site.id}/verify-checkin", json={"location_unavailable": True}
        ).status_code, 400)
        self.assertEqual(self.client.post(f"/visit/{self.site.id}/checkin").status_code, 403)
        self.client.post(f"/api/visit/{self.site.id}/verify-checkin", json={"latitude": 0, "longitude": 0.01})
        with self.client.session_transaction() as stored:
            verification = stored["visit_verification"]
            verification["verified_at"] = (datetime.now(timezone.utc) - timedelta(minutes=11)).isoformat()
            stored["visit_verification"] = verification
        self.assertEqual(self.client.get(f"/visit/{self.site.id}/camera").status_code, 403)
        self.assertEqual(self.client.post(f"/visit/{self.site.id}/checkin").status_code, 403)
        self.assertEqual(Attendance.query.count(), 0)

    def test_daily_cutoff_is_exact_india_time_and_catches_up_idempotently(self):
        visit = self._start_location_visit()
        visit.checkin_time = datetime(2026, 10, 20, 5, tzinfo=timezone.utc)
        visit.last_location_time = datetime(2026, 10, 20, 18, 28, tzinfo=timezone.utc)
        visit.last_location_lon = 0.04
        db.session.commit()
        cutoff = datetime(2026, 10, 20, 18, 29, tzinfo=timezone.utc)
        self.assertEqual(daily_cutoff(visit.checkin_time), cutoff)
        self.assertEqual(close_overdue_visits(cutoff - timedelta(microseconds=1)), 0)
        self.assertEqual(close_overdue_visits(cutoff), 1)
        db.session.refresh(visit)
        self.assertEqual(visit.status, "report_pending")
        self.assertEqual(as_utc(visit.checkout_time), cutoff)
        self.assertEqual(visit.checkout_lon, 0.04)
        self.assertEqual(visit.checkout_location_source, "last_known")
        self.assertEqual(visit.checkout_location_reason, "daily_cutoff")
        self.assertTrue(visit.auto_closed)
        self.assertEqual(close_overdue_visits(cutoff + timedelta(days=3)), 0)
        self.assertEqual(AuditLog.query.filter_by(action="site_auto_checkout").count(), 1)
        self.assertEqual(SiteReport.query.count(), 0)

    def test_scheduler_catches_up_offline_visits_and_report_keeps_stop_details(self):
        visit = self._start_location_visit()
        visit.checkin_time = datetime(2026, 10, 19, 6, tzinfo=timezone.utc)
        visit.last_location_time = visit.checkin_time
        db.session.commit()
        result = self.app.test_cli_runner().invoke(args=["close-overdue-visits"])
        self.assertEqual(result.exit_code, 0, result.output)
        self.assertIn("stopped 1 visit", result.output)
        db.session.refresh(visit)
        stopped = (visit.checkout_time, visit.checkout_lat, visit.checkout_lon, visit.checkout_location_time)
        self.assertEqual(as_utc(visit.checkout_time), datetime(2026, 10, 19, 18, 29, tzinfo=timezone.utc))
        self.assertEqual(self.client.get(f"/visit/{visit.id}/report").status_code, 200)
        response = self.client.post(f"/visit/{visit.id}/report", data={
            "activity_detail": "Work before cutoff", "billing": "billable", "engineer_name": "Engineer",
        })
        self.assertEqual(response.status_code, 302)
        db.session.refresh(visit)
        self.assertEqual(visit.status, "completed")
        self.assertEqual((visit.checkout_time, visit.checkout_lat, visit.checkout_lon, visit.checkout_location_time), stopped)
        self.assertTrue(visit.auto_closed)
        self.assertEqual(self.client.post(f"/visit/{visit.id}/report", data={}).status_code, 403)
        self.assertEqual(SiteReport.query.count(), 1)

    def test_request_catches_overdue_visit_and_admin_sees_distances_and_provenance(self):
        visit = self._start_location_visit()
        visit.checkin_time = datetime(2026, 10, 20, 5, tzinfo=timezone.utc)
        visit.last_location_time = visit.checkin_time
        db.session.commit()
        response = self.client.get("/employee")
        self.assertIn(b"Auto-stopped at the daily 23:59 India cutoff", response.data)
        self.assertIn(b"Complete report", response.data)
        self.assertNotIn(b"Onsite now", response.data)
        self.client.post("/logout")
        self.client.post("/login", data={"email": self.admin.email, "password": "admin-password-123"})
        response = self.client.get("/admin/attendance")
        self.assertEqual(response.status_code, 200)
        for text in (b"1.112 km from assigned site", b"Last captured location", b"Daily Cutoff",
                     b"Report Pending", b"0.00000, 0.01000", b"Auto-stopped (23:59 India)"):
            self.assertIn(text, response.data)
        live = self.client.get("/api/admin/live-status").json
        self.assertEqual(live["onsite_count"], 0)
        self.assertEqual(live["employees"][0]["status"], "report-pending")
        self.assertIn(b"Stopped visits", self.client.get("/admin/metrics/pending-reports").data)

    def test_location_and_pending_reports_are_owner_scoped(self):
        visit = self._start_location_visit()
        other = User(name="Other", email="other@example.com", mobile="3")
        other.set_password("other-password-123")
        db.session.add(other)
        db.session.commit()
        # Isolate Flask-Login's request user cache from the outer test app context.
        with self.app.app_context():
            client = self.app.test_client()
            client.post("/login", data={"email": other.email, "password": "other-password-123"})
            for path in ("location", "verify-checkout"):
                response = client.post(f"/api/visit/{visit.id}/{path}", json={"latitude": 0, "longitude": 0})
                self.assertEqual(response.status_code, 404)
            self.assertEqual(client.get(f"/visit/{visit.id}/report").status_code, 404)
            self.assertEqual(client.get("/admin/attendance").status_code, 403)
        self.assertEqual(self.client.get(f"/visit/{visit.id}/report").status_code, 403)
        self.assertEqual(self.client.post(
            f"/api/visit/{self.site.id}/verify-checkin", json={"latitude": 0, "longitude": 0}
        ).status_code, 409)

    def test_daily_cutoff_blocks_new_start_until_local_midnight(self):
        self._login_employee()
        cutoff = datetime(2026, 10, 26, 18, 29, tzinfo=timezone.utc)
        endpoint = f"/api/visit/{self.site.id}/verify-checkin"
        with patch("geosite.routes.utcnow", return_value=cutoff):
            self.assertEqual(self.client.post(endpoint, json={"latitude": 0, "longitude": 0}).status_code, 409)
        with patch("geosite.routes.utcnow", return_value=cutoff + timedelta(minutes=1)):
            self.assertEqual(self.client.post(endpoint, json={"latitude": 0, "longitude": 0}).status_code, 200)

    def test_pending_report_does_not_block_new_visit_or_clear_its_location_capture(self):
        visit = self._start_location_visit()
        self.client.post(f"/api/visit/{visit.id}/verify-checkout", json={
            "location_unavailable": True, "reason": "permission_denied",
        })
        stopped_at = visit.checkout_time
        self.assertEqual(self.client.post(
            f"/api/visit/{self.site.id}/verify-checkin", json={"latitude": 0, "longitude": 0.02}
        ).status_code, 200)
        response = self.client.post(f"/visit/{visit.id}/report", data={
            "activity_detail": "Earlier visit", "billing": "non_billable", "engineer_name": "Engineer",
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.client.get(f"/visit/{self.site.id}/camera").status_code, 200)
        self.client.post(f"/visit/{self.site.id}/checkin", data={"selfie": self._image_data()})
        self.assertEqual(Attendance.query.filter_by(status="onsite").count(), 1)
        self.assertEqual(Attendance.query.filter_by(status="completed").count(), 1)
        db.session.refresh(visit)
        self.assertEqual(visit.checkout_time, stopped_at)

    def _login_admin(self):
        self.client.post("/logout")
        return self.client.post("/login", data={
            "email": self.admin.email, "password": "admin-password-123",
        })

    def test_project_list_crud_links_sites_and_is_admin_only(self):
        response = self.client.get("/admin/projects")
        self.assertEqual(response.status_code, 302)
        self._login_admin()
        response = self.client.get("/admin/projects")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Project list", response.data)
        self.assertIn(b"Project", self.client.get("/admin/sites").data)
        self.client.post("/admin/projects", data={"name": "Missing number"})
        self.assertIsNone(Project.query.filter_by(name="Missing number").first())

        self.client.post("/admin/projects", data={
            "project_number": " prj-001 ", "name": "  New Project  ",
        })
        project = Project.query.filter_by(name="New Project").one()
        self.assertEqual(project.project_number, "PRJ-001")
        self.client.post("/admin/projects", data={
            "project_number": "prj-001", "name": "Duplicate number",
        })
        self.assertIsNone(Project.query.filter_by(name="Duplicate number").first())
        site_payload = self._site_payload(project_id=str(project.id))
        response = self.client.post("/admin/sites", data=site_payload, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        site = AssignedSite.query.filter_by(site_name="South Site").one()
        self.assertEqual(site.project_id, project.id)
        self.assertEqual(site.project_name, "New Project")
        self.assertEqual(site.report_company, REPORT_COMPANY_NAME)
        self.assertIn(b"Blackpurple automation", response.data)

        self.client.post(f"/admin/projects/{project.id}/edit", data={
            "project_number": "PRJ-101", "name": "Renamed Project",
        })
        self.assertEqual(project.project_number, "PRJ-101")
        db.session.refresh(site)
        self.assertEqual(site.project_name, "Renamed Project")
        self.assertIn(b"PRJ-101 - Renamed Project", self.client.get("/admin/sites").data)
        response = self.client.post(
            f"/admin/projects/{project.id}/delete", data={}, follow_redirects=True,
        )
        self.assertIn(b"cannot be deleted", response.data)
        self.assertIsNotNone(db.session.get(Project, project.id))

        self.client.post("/logout")
        self._login_employee()
        self.assertEqual(self.client.get("/admin/projects").status_code, 403)
        self.assertEqual(
            self.client.post("/admin/projects", data={"name": "Unauthorized"}).status_code,
            403,
        )

    def _site_payload(self, **updates):
        return {
            "employee_id": str(self.employee.id), "site_name": "South Site",
            "customer_name": "South Customer", "address": "2 Example Road",
            "latitude": "19", "longitude": "72", "allowed_radius": "250",
            "project_id": str(self.project.id),
            "report_company": "Reporting Co", "commissioning_manager": "Commissioning Lead",
            "project_manager": "Project Lead", "report_remarks": "First\nSecond",
            **updates,
        }

    def _completed_display_visit(self, stopped_at):
        visit = Attendance(
            employee_id=self.employee.id, site_id=self.site.id, status="completed",
            checkin_time=stopped_at - timedelta(hours=1), checkout_time=stopped_at,
            checkin_lat=0, checkin_lon=0.01, checkout_lat=0, checkout_lon=0.02,
            checkin_selfie="missing.png", checkin_distance_m=1111.949,
            checkout_distance_m=2223.899, checkout_location_source="current",
            checkout_location_time=stopped_at,
        )
        db.session.add(visit)
        db.session.flush()
        db.session.add(SiteReport(
            attendance=visit, work_performed="Display test", is_billable=True,
            activities_completed="", issues_found="", recommendations="", materials_used="",
            remarks="", engineer_name="Engineer", customer_representative_name="",
            employee_signature="", customer_signature="",
        ))
        db.session.commit()
        return visit

    def test_pdf_and_attendance_show_both_distances_in_km_and_ist(self):
        visit = self._completed_display_visit(datetime(2026, 10, 26, 20, tzinfo=timezone.utc))
        self._login_admin()
        response = self.client.get(f"/reports/{visit.id}/pdf")
        self.assertEqual(response.status_code, 200)
        text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.data)).pages)
        for value in ("1.112 km", "2.224 km", "27 Oct 2026 00:30 IST", "27 Oct 2026 01:30 IST"):
            self.assertIn(value, text)
        self.assertNotIn(" UTC", text)
        response.close()
        history = self.client.get("/admin/attendance")
        for value in (b"1.112 km from assigned site", b"2.224 km from assigned site",
                      b"27 Oct 2026 00:30 IST", b"27 Oct 2026 01:30 IST"):
            self.assertIn(value, history.data)
        self.assertIn(b"27 Oct 2026 00:30 IST", self.client.get("/reports").data)
        # Unknown distances remain unknown; a captured zero remains zero.
        visit.checkin_distance_m, visit.checkout_distance_m = 0, None
        db.session.commit()
        response = self.client.get(f"/reports/{visit.id}/pdf")
        text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.data)).pages)
        self.assertIn("0.000 km", text)
        self.assertIn("Not recorded", text)
        response.close()
        self.assertIn(b"0.000 km from assigned site", self.client.get("/admin/attendance").data)

    def test_employee_pdf_downloads_are_denied_but_history_and_excel_remain(self):
        visit = self._completed_display_visit(datetime(2026, 10, 26, 12, tzinfo=timezone.utc))
        self._login_employee()
        for path in ("/employee", "/reports"):
            page = self.client.get(path)
            self.assertEqual(page.status_code, 200)
            self.assertNotIn(b">PDF<", page.data)
            self.assertNotIn(b"Download PDF", page.data)
            self.assertNotIn(b"download-combined-pdf", page.data)
        self.assertEqual(self.client.get(f"/reports/{visit.id}/pdf").status_code, 403)
        self.assertEqual(self.client.get("/reports/combined.pdf", query_string={
            "start_date": "2026-10-26", "end_date": "2026-10-26",
        }).status_code, 403)
        response = self.client.get("/reports/activity.xlsx", query_string={
            "site_id": self.site.id, "start_date": "2026-10-26", "end_date": "2026-10-26",
        })
        self.assertEqual(response.status_code, 200)
        response.close()
        self._login_admin()
        self.assertIn(b"Download PDF", self.client.get("/reports").data)
        self.assertIn(b"download-combined-pdf", self.client.get("/reports").data)

    def test_po_fields_only_appear_on_invoices_not_site_catalog(self):
        self._login_admin()
        for path in ("/admin/sites", f"/admin/sites/{self.site.id}/edit"):
            page = self.client.get(path)
            self.assertNotIn(b'name="po_number"', page.data)
            self.assertNotIn(b'name="po_date"', page.data)
            self.assertNotIn(b'name="employee_id"', page.data)
        page = self.client.get("/admin/invoices")
        self.assertIn(b'name="po_number"', page.data)
        self.assertIn(b'name="po_date"', page.data)

    def test_admin_today_metrics_and_combined_pdf_use_india_date_boundaries(self):
        stopped_at = datetime(2026, 10, 26, 20, tzinfo=timezone.utc)
        visit = self._completed_display_visit(stopped_at)
        self._login_admin()
        with patch("geosite.routes.utcnow", return_value=stopped_at + timedelta(minutes=10)):
            page = self.client.get("/admin/metrics/completed")
            self.assertIn(b"27 Oct 2026 01:30 IST", page.data)
            self.assertIn(b"Completed", self.client.get("/admin").data)
            self.assertEqual(self.client.get("/api/admin/live-status").json["employees"][0]["status"], "completed")
        included = self.client.get("/reports/combined.pdf", query_string={
            "start_date": "2026-10-27", "end_date": "2026-10-27",
        })
        self.assertEqual(included.status_code, 200)
        included.close()
        excluded = self.client.get("/reports/combined.pdf", query_string={
            "start_date": "2026-10-26", "end_date": "2026-10-26",
        })
        self.assertEqual(excluded.status_code, 302)
        self.assertEqual(visit.checkout_time, stopped_at.replace(tzinfo=None))

    def test_admin_live_checkin_times_match_ist_in_html_and_polling(self):
        self._login_admin()
        now = datetime.now(timezone.utc)
        visit = Attendance(
            employee_id=self.employee.id, site_id=self.site.id, status="onsite",
            checkin_time=now, checkin_lat=0, checkin_lon=0, checkin_selfie="missing.png",
        )
        db.session.add(visit)
        db.session.commit()
        expected = now.astimezone(INDIA_TIME).strftime("%H:%M IST")
        self.assertIn(expected.encode(), self.client.get("/admin").data)
        self.assertEqual(self.client.get("/api/admin/live-status").json["employees"][0]["checkin_time"], expected)
        self.assertIn(b"IST", self.client.get("/admin/metrics/onsite").data)

    def test_completed_metric_uses_exact_ist_midnight_boundaries(self):
        midnight = datetime(2026, 10, 26, 18, 30, tzinfo=timezone.utc)
        for stopped_at in (
            midnight - timedelta(microseconds=1),
            midnight,
            midnight + timedelta(days=1) - timedelta(microseconds=1),
            midnight + timedelta(days=1),
        ):
            self._completed_display_visit(stopped_at)
        self._login_admin()
        with patch("geosite.routes.utcnow", return_value=midnight + timedelta(hours=1)):
            with patch("geosite.routes.render_template", return_value="dashboard") as render:
                response = self.client.get("/admin")
                self.assertEqual(response.status_code, 200)
                values = render.call_args.kwargs
                self.assertNotIn("completed_today", values)
                self.assertNotIn("chart_labels", values)
                self.assertNotIn("chart_values", values)
            with patch("geosite.routes.render_template", return_value="metric") as render:
                self.client.get("/admin/metrics/completed")
                self.assertEqual(len(render.call_args.kwargs["rows"]), 2)

    def test_login_password_toggle_is_accessible_and_not_a_submit_button(self):
        page = self.client.get("/login")
        self.assertIn(b'id="password" type="password"', page.data)
        self.assertIn(b'id="toggle-password" type="button"', page.data)
        self.assertIn(b'aria-label="Show password" aria-pressed="false"', page.data)

    def test_admin_assigns_multiple_sites_without_replacing_existing_assignments(self):
        self._login_admin()
        response = self.client.post("/admin/sites", data=self._site_payload(), follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        second = AssignedSite.query.filter_by(site_name="South Site").one()
        self.assertEqual(SiteAssignment.query.filter_by(employee_id=self.employee.id, is_active=True).count(), 1)
        self.client.post("/admin/assignments", data={
            "employee_id": self.employee.id, "site_ids": [self.site.id, second.id],
        })
        self.assertEqual(SiteAssignment.query.filter_by(employee_id=self.employee.id, is_active=True).count(), 2)
        db.session.refresh(self.site)
        self.assertEqual(self.site.site_name, "North Site")
        self.assertEqual(self.site.latitude, 0)
        self.assertTrue(self.site.is_active)
        self.assertIn(b"North Site, South Site", self.client.get("/admin").data)
        live = self.client.get("/api/admin/live-status").json["employees"][0]
        self.assertEqual(live["site"], "North Site, South Site")
        self.client.post("/logout")
        dashboard = self._login_employee()
        self.assertIn(b"North Site", dashboard.data)
        self.assertIn(b"South Site", dashboard.data)
        second = AssignedSite.query.filter_by(site_name="South Site").one()
        self.assertEqual(self.client.post(f"/visit/start/{second.id}").status_code, 302)
        self.assertEqual(self.client.post(
            f"/api/visit/{second.id}/verify-checkin", json={"latitude": 19, "longitude": 72}
        ).status_code, 200)
        self.client.post(f"/visit/{second.id}/checkin", data={"selfie": self._image_data()})
        self.assertEqual(Attendance.query.one().site_id, second.id)
        self.assertEqual(self.client.post(
            f"/api/visit/{self.site.id}/verify-checkin", json={"latitude": 0, "longitude": 0}
        ).status_code, 409)

    def test_edit_assignment_prefills_and_updates_only_selected_site(self):
        self._login_admin()
        self.client.post("/admin/sites", data=self._site_payload())
        second = AssignedSite.query.filter_by(site_name="South Site").one()
        url = f"/admin/sites/{second.id}/edit"
        form = self.client.get(url)
        self.assertEqual(form.status_code, 200)
        for value in (b'value="South Site"', b'value="19.0"', b'Project Two',
                      b'Blackpurple automation', b'First\nSecond', b'Save changes'):
            self.assertIn(value, form.data)
        updated_project = Project(project_number="PRJ-004", name="Updated Project")
        db.session.add(updated_project)
        db.session.commit()
        response = self.client.post(url, data=self._site_payload(
            site_name="South Updated", customer_name="Updated Customer", address="New Address",
            latitude="20", longitude="73", allowed_radius="500", project_id=str(updated_project.id),
            report_company="Updated Company",
        ), follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        db.session.refresh(second)
        self.assertEqual((second.site_name, second.customer_name, second.address),
                         ("South Updated", "Updated Customer", "New Address"))
        self.assertEqual((second.latitude, second.longitude, second.allowed_radius), (20, 73, 500))
        self.assertEqual(second.project_name, "Updated Project")
        self.assertEqual(second.report_company, REPORT_COMPANY_NAME)
        self.assertEqual(AssignedSite.query.count(), 2)
        db.session.refresh(self.site)
        self.assertEqual(self.site.site_name, "North Site")
        self.assertTrue(self.site.is_active)
        self.assertEqual(AuditLog.query.filter_by(action="site_updated").count(), 1)

    def test_assignment_edits_are_blocked_onsite_but_other_sites_can_be_added(self):
        visit = self._start_location_visit()
        with self.app.app_context():
            client = self.app.test_client()
            client.post("/login", data={"email": self.admin.email, "password": "admin-password-123"})
            url = f"/admin/sites/{self.site.id}/edit"
            for method in ("get", "post"):
                response = getattr(client, method)(url, data=self._site_payload() if method == "post" else None,
                                                   follow_redirects=True)
                self.assertIn(b"cannot be edited during an active visit", response.data)
            client.post("/admin/sites", data=self._site_payload())
        db.session.refresh(self.site)
        self.assertEqual(self.site.site_name, "North Site")
        self.assertEqual(AssignedSite.query.filter_by(is_active=True).count(), 2)
        self.assertEqual(SiteAssignment.query.filter_by(employee_id=self.employee.id, is_active=True).count(), 1)
        db.session.refresh(visit)
        self.assertEqual(visit.status, "onsite")

    def test_historical_assignment_can_be_edited_without_reassigning_engineer(self):
        visit = self._start_location_visit()
        self.client.post(f"/api/visit/{visit.id}/verify-checkout", json={"latitude": 0, "longitude": 0.02})
        snapshots = (visit.site_latitude, visit.site_longitude, visit.checkin_distance_m, visit.checkout_distance_m)
        self._login_admin()
        other = User(name="Other", email="other@example.com", mobile="3", password_hash="test")
        db.session.add(other)
        db.session.commit()
        url = f"/admin/sites/{self.site.id}/edit"
        response = self.client.post(url, data=self._site_payload(employee_id=str(other.id)), follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        db.session.refresh(self.site)
        self.assertEqual(SiteAssignment.query.filter_by(site_id=self.site.id).one().employee_id, self.employee.id)
        db.session.refresh(self.site)
        self.assertEqual(self.site.site_name, "South Site")
        self.assertEqual((self.site.latitude, self.site.longitude), (19, 72))
        self.assertEqual(AssignedSite.query.count(), 1)
        db.session.refresh(visit)
        self.assertEqual((visit.site_latitude, visit.site_longitude, visit.checkin_distance_m, visit.checkout_distance_m), snapshots)
        self.assertEqual(visit.employee_id, self.employee.id)

    def test_site_edit_cannot_change_engineer_access(self):
        self._login_admin()
        other = User(name="Other", email="other@example.com", mobile="3", password_hash="test")
        db.session.add(other)
        db.session.commit()
        self.client.post(f"/admin/sites/{self.site.id}/edit", data=self._site_payload(employee_id=str(other.id)))
        db.session.refresh(self.site)
        self.assertEqual(SiteAssignment.query.filter_by(site_id=self.site.id).one().employee_id, self.employee.id)
        self.assertEqual(AssignedSite.query.count(), 1)

    def test_invoice_history_retains_engineer_when_shared_site_is_edited(self):
        self._login_admin()
        self.client.post("/admin/invoices", data={
            "site_id": self.site.id, "employee_id": self.employee.id, "customer_name": "Example Customer",
            "description": "Service", "start_date": "2026-10-10", "end_date": "2026-10-10",
            "billable_days": "1", "daily_rate": "100", "currency": "INR",
        })
        self.assertEqual(Invoice.query.count(), 1)
        other = User(name="Other", email="other@example.com", mobile="3", password_hash="test")
        db.session.add(other)
        db.session.commit()
        url = f"/admin/sites/{self.site.id}/edit"
        response = self.client.post(url, data=self._site_payload(employee_id=str(other.id)), follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.client.post(url, data=self._site_payload())
        db.session.refresh(self.site)
        self.assertEqual(self.site.site_name, "South Site")
        self.assertEqual(SiteAssignment.query.filter_by(site_id=self.site.id).one().employee_id, self.employee.id)
        self.assertEqual(Invoice.query.one().customer_name, "Example Customer")
        self.assertEqual(Invoice.query.one().site_id, self.site.id)
        self.assertEqual(Invoice.query.one().employee_id, self.employee.id)

    def test_invalid_assignment_edit_keeps_saved_values_and_preserves_form_input(self):
        self._login_admin()
        url = f"/admin/sites/{self.site.id}/edit"
        for payload in (self._site_payload(latitude="91"), self._site_payload(longitude="nan"),
                        self._site_payload(allowed_radius="0"),
                        self._site_payload(site_name="x" * 161), self._site_payload(project_id="invalid")):
            response = self.client.post(url, data=payload)
            self.assertEqual(response.status_code, 200)
            self.assertIn(b"Enter valid site details", response.data)
            self.assertIn(b"2 Example Road", response.data)
            db.session.refresh(self.site)
            self.assertEqual(self.site.site_name, "North Site")
            self.assertEqual(self.site.latitude, 0)
        self.assertEqual(AssignedSite.query.count(), 1)
        self.assertEqual(AuditLog.query.filter_by(action="site_updated").count(), 0)

    def test_employee_cannot_create_or_edit_assignments(self):
        self._login_employee()
        self.assertEqual(self.client.post("/admin/sites", data=self._site_payload()).status_code, 403)
        url = f"/admin/sites/{self.site.id}/edit"
        self.assertEqual(self.client.get(url).status_code, 403)
        self.assertEqual(self.client.post(url, data=self._site_payload()).status_code, 403)
        self.assertEqual(AssignedSite.query.count(), 1)
        self._login_admin()
        self.assertEqual(self.client.get("/admin/sites/999/edit").status_code, 404)

    def test_assignment_engineers_render_as_wrapping_searchable_radio_list(self):
        self._login_admin()
        page = self.client.get("/admin/assignments", query_string={"employee_id": self.employee.id})
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'id="assignment-search"', page.data)
        self.assertIn(b'class="assignment-engineers"', page.data)
        self.assertIn(b'type="radio" name="employee_id"', page.data)
        self.assertIn(b"engineer@example.com", page.data)
        self.assertIn(b"assignments.js", page.data)
        self.assertIn(b"Load assignments", page.data)
        self.assertIn(b"Save assignments", page.data)
        self.assertNotIn(b'id="assignment-engineer"', page.data)

    def test_shared_catalog_site_can_be_assigned_to_multiple_engineers_without_copying(self):
        self._login_admin()
        other = User(name="Other", email="other@example.com", mobile="3")
        other.set_password("other-password-123")
        db.session.add(other)
        db.session.commit()
        for _ in range(2):
            response = self.client.post("/admin/assignments", data={
                "employee_id": other.id, "site_ids": [self.site.id],
            }, follow_redirects=True)
            self.assertIn(b"Engineer site assignments saved", response.data)
        self.assertEqual(AssignedSite.query.count(), 1)
        self.assertEqual(SiteAssignment.query.filter_by(site_id=self.site.id, is_active=True).count(), 2)
        self.assertEqual(SiteAssignment.query.filter_by(employee_id=other.id, site_id=self.site.id).count(), 1)
        self.assertIn(b"Engineer, Other", self.client.get("/admin/sites").data)
        with self.app.app_context():
            client = self.app.test_client()
            client.post("/login", data={"email": other.email, "password": "other-password-123"})
            self.assertIn(b"North Site", client.get("/employee").data)
            self.assertEqual(client.post(
                f"/api/visit/{self.site.id}/verify-checkin", json={"latitude": 0, "longitude": 0}
            ).status_code, 200)
            client.post(f"/visit/{self.site.id}/checkin", data={"selfie": self._image_data()})
        visit = Attendance.query.one()
        self.assertEqual(visit.employee_id, other.id)
        self.assertEqual(visit.site_id, self.site.id)

    def test_unassign_and_reassign_preserve_history_without_new_site_or_membership(self):
        visit = self._completed_display_visit(datetime(2026, 10, 26, 12, tzinfo=timezone.utc))
        self._login_admin()
        self.client.post("/admin/assignments", data={"employee_id": self.employee.id})
        self.assertFalse(SiteAssignment.query.one().is_active)
        self.assertEqual(Attendance.query.one().id, visit.id)
        self.client.post("/logout")
        self._login_employee()
        self.assertNotIn(f'<option value="{self.site.id}">'.encode(), self.client.get("/employee").data)
        self.assertEqual(self.client.post(
            f"/api/visit/{self.site.id}/verify-checkin", json={"latitude": 0, "longitude": 0}
        ).status_code, 404)
        export = self.client.get("/reports/activity.xlsx", query_string={
            "site_id": self.site.id, "start_date": "2026-10-26", "end_date": "2026-10-26",
        })
        self.assertEqual(export.status_code, 200)
        self.assertIn("Display test", str(list(load_workbook(io.BytesIO(export.data)).active.values)))
        self._login_admin()
        self.client.post("/admin/assignments", data={"employee_id": self.employee.id, "site_ids": [self.site.id]})
        self.assertTrue(SiteAssignment.query.one().is_active)
        self.assertEqual(AssignedSite.query.count(), 1)

    def test_unassign_active_visit_is_rejected_atomically(self):
        self._start_location_visit()
        second = AssignedSite(site_name="Extra", customer_name="Customer", address="Address",
                              latitude=1, longitude=1, allowed_radius=100)
        db.session.add(second)
        db.session.commit()
        with self.app.app_context():
            client = self.app.test_client()
            client.post("/login", data={"email": self.admin.email, "password": "admin-password-123"})
            response = client.post("/admin/assignments", data={
                "employee_id": self.employee.id, "site_ids": [second.id],
            }, follow_redirects=True)
            self.assertIn(b"No assignments were changed", response.data)
        db.session.expire_all()
        self.assertEqual(SiteAssignment.query.count(), 1)
        self.assertTrue(SiteAssignment.query.one().is_active)

    def test_assignment_input_and_roles_are_validated(self):
        self._login_employee()
        self.assertEqual(self.client.get("/admin/assignments").status_code, 403)
        self.assertEqual(self.client.post("/admin/assignments", data={
            "employee_id": self.employee.id, "site_ids": [self.site.id],
        }).status_code, 403)
        self._login_admin()
        for payload, status in (
            ({}, 400), ({"employee_id": "bad"}, 400),
            ({"employee_id": self.admin.id}, 404),
            ({"employee_id": self.employee.id, "site_ids": ["bad"]}, 400),
            ({"employee_id": self.employee.id, "site_ids": ["999"]}, 400),
        ):
            self.assertEqual(self.client.post("/admin/assignments", data=payload).status_code, status)
        self.assertTrue(SiteAssignment.query.one().is_active)
        self.assertEqual(SiteAssignment.query.count(), 1)
        self.site.is_active = False
        db.session.commit()
        self.assertEqual(self.client.post("/admin/assignments", data={
            "employee_id": self.employee.id, "site_ids": [self.site.id],
        }).status_code, 400)
        self.app.config["WTF_CSRF_ENABLED"] = True
        self.assertEqual(self.client.post("/admin/assignments", data={"employee_id": self.employee.id}).status_code, 400)

    def test_shared_site_excel_isolates_each_engineers_activities_leaves_and_header(self):
        visit = self._completed_display_visit(datetime(2026, 10, 26, 12, tzinfo=timezone.utc))
        other = User(name="Other", email="other@example.com", mobile="3", password_hash="test")
        db.session.add(other)
        db.session.flush()
        db.session.add(SiteAssignment(employee=other, site=self.site))
        other_visit = Attendance(
            employee_id=other.id, site_id=self.site.id, status="completed",
            checkin_lat=0, checkin_lon=0, checkin_selfie="missing.png",
            checkout_time=datetime(2026, 10, 26, 13, tzinfo=timezone.utc),
        )
        db.session.add(other_visit)
        db.session.flush()
        db.session.add(SiteReport(
            attendance=other_visit, work_performed="Other private work", is_billable=True,
            activities_completed="", issues_found="", recommendations="", materials_used="",
            remarks="", engineer_name="Other", customer_representative_name="",
            employee_signature="", customer_signature="",
        ))
        db.session.add(LeaveRequest(
            employee_id=other.id, start_date=date(2026, 10, 26), end_date=date(2026, 10, 26),
            period="full", reason="Other private leave", status="approved",
        ))
        db.session.commit()
        self._login_employee()
        params = {"site_id": self.site.id, "employee_id": other.id,
                  "start_date": "2026-10-26", "end_date": "2026-10-26"}
        response = self.client.get("/reports/activity.xlsx", query_string=params)
        sheet = load_workbook(io.BytesIO(response.data)).active
        values = str(list(sheet.values))
        self.assertIn("ENGINEER : Engineer", values)
        self.assertIn("Display test", values)
        self.assertNotIn("Other private", values)
        self.assertNotIn("PO NO", values)
        self.assertNotIn("PO DATE", values)
        self._login_admin()
        response = self.client.get("/reports/activity.xlsx", query_string=params)
        values = str(list(load_workbook(io.BytesIO(response.data)).active.values))
        self.assertIn("ENGINEER : Other", values)
        self.assertIn("Other private work", values)
        self.assertIn("Other private leave", values)
        self.assertNotIn("Display test", values)

    def test_invoice_po_is_manual_optional_validated_and_not_taken_from_site(self):
        self._login_admin()
        payload = {
            "site_id": self.site.id, "employee_id": self.employee.id,
            "customer_name": "Customer", "description": "Service",
            "start_date": "2026-10-26", "end_date": "2026-10-26", "billable_days": "1",
            "daily_rate": "100", "currency": "INR", "po_number": "MANUAL-42", "po_date": "2026-10-01",
        }
        self.client.post("/admin/invoices", data=payload)
        invoice = Invoice.query.one()
        self.assertEqual(invoice.po_number, "MANUAL-42")
        self.assertEqual(invoice.po_date, date(2026, 10, 1))
        response = self.client.get(f"/admin/invoices/{invoice.id}/pdf")
        text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.data)).pages)
        self.assertIn("MANUAL-42", text)
        self.assertIn("2026-10-01", text)
        response.close()
        for updates in ({"po_date": "bad"}, {"po_number": "x" * 121},
                        {"employee_id": self.admin.id}, {"employee_id": ""}, {"site_id": 999}):
            self.client.post("/admin/invoices", data={**payload, **updates})
            self.assertEqual(Invoice.query.count(), 1)
        self.client.post("/admin/invoices", data={**payload, "po_number": "", "po_date": ""})
        self.assertEqual(Invoice.query.count(), 2)
        blank = Invoice.query.order_by(Invoice.id.desc()).first()
        self.assertEqual(blank.po_number, "")
        self.assertIsNone(blank.po_date)
        other = User(name="Other", email="other@example.com", mobile="3", password_hash="test")
        db.session.add(other)
        db.session.commit()
        self.client.post("/admin/invoices", data={**payload, "employee_id": other.id})
        self.assertEqual(Invoice.query.count(), 2)
        self.client.post("/admin/assignments", data={"employee_id": other.id, "site_ids": [self.site.id]})
        self.client.post("/admin/invoices", data={**payload, "employee_id": other.id, "po_number": "OTHER-PO"})
        latest = Invoice.query.order_by(Invoice.id.desc()).first()
        self.assertEqual(latest.employee_id, other.id)
        self.assertEqual(latest.site_id, self.site.id)
        self.assertEqual(latest.po_number, "OTHER-PO")

    def test_deleting_engineer_without_visit_history_keeps_shared_catalog_site(self):
        other = User(name="Other", email="other@example.com", mobile="3", password_hash="test")
        db.session.add(other)
        db.session.flush()
        db.session.add(SiteAssignment(employee=other, site=self.site))
        db.session.commit()
        other_id = other.id
        self._login_admin()
        self.client.post(f"/admin/employees/{other_id}/delete")
        self.assertIsNone(db.session.get(User, other_id))
        self.assertIsNotNone(db.session.get(AssignedSite, self.site.id))
        self.assertEqual(SiteAssignment.query.filter_by(site_id=self.site.id).count(), 1)
        self.assertEqual(SiteAssignment.query.one().employee_id, self.employee.id)

    def test_admin_invoices_use_days_times_daily_rate(self):
        self.client.post(
            "/login",
            data={"email": self.admin.email, "password": "admin-password-123"},
        )
        response = self.client.post(
            "/admin/invoices",
            data={
                "site_id": self.site.id,
                "employee_id": self.employee.id,
                "customer_name": "Example Customer",
                "description": "Engineering service",
                "start_date": "2026-10-10",
                "end_date": "2026-10-12",
                "billable_days": "2",
                "daily_rate": "123.45",
                "currency": "INR",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        invoice = Invoice.query.one()
        self.assertEqual(str(invoice.total_amount), "246.90")
        pdf = self.client.get(f"/admin/invoices/{invoice.id}/pdf")
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(len(PdfReader(io.BytesIO(pdf.data)).pages), 1)

    def test_dashboard_cards_open_their_detail_pages(self):
        self.client.post(
            "/login",
            data={"email": self.admin.email, "password": "admin-password-123"},
        )
        dashboard = self.client.get("/admin")
        self.assertNotIn(b"completion-chart", dashboard.data)
        self.assertNotIn(b"chart.js", dashboard.data)
        self.assertIn(b'id="live-map"', dashboard.data)
        for label in (b"Completed today", b"Reports submitted", b"Reports pending"):
            self.assertNotIn(label, dashboard.data)
        for metric in ("employees", "sites", "onsite"):
            self.assertIn(f"/admin/metrics/{metric}".encode(), dashboard.data)
        for metric in ("completed", "reports", "pending-reports"):
            self.assertNotIn(f"/admin/metrics/{metric}".encode(), dashboard.data)
        self.assertIn(b"Pending leave requests", dashboard.data)
        self.assertIn(b"On leave today", dashboard.data)
        for metric in ("employees", "sites", "onsite", "completed", "reports", "pending-reports"):
            with self.subTest(metric=metric):
                self.assertEqual(self.client.get(f"/admin/metrics/{metric}").status_code, 200)

    def test_admin_pages_render(self):
        self.client.post(
            "/login",
            data={"email": self.admin.email, "password": "admin-password-123"},
        )
        for path in (
            "/admin",
            "/admin/employees",
            "/admin/sites",
            "/admin/invoices",
            "/reports",
            "/admin/attendance",
        ):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200)

    def test_admin_creates_and_edits_optional_employee_details(self):
        self._login_admin()
        payload = {
            "name": "New Engineer", "email": "new@example.com",
            "mobile": "1234567890", "designation": "Engineer",
            "password": "new-password-123",
            "address": "  15 Field Road\nMumbai  ",
            "emergency_contact_no": "+91 9876543210",
            "blood_group": "o+", "laptop_serial_no": "LAP-100",
            "employee_code": " eng-101 ",
            "birthday": "1996-02-29",
        }
        response = self.client.post("/admin/employees", data=payload, follow_redirects=True)
        self.assertIn(b"Employee created.", response.data)
        employee = User.query.filter_by(email=payload["email"]).one()
        self.assertEqual(employee.address, "15 Field Road\nMumbai")
        self.assertEqual(employee.blood_group, "O+")
        self.assertEqual(employee.employee_code, "ENG-101")
        self.assertEqual(employee.emergency_contact_no, payload["emergency_contact_no"])
        self.assertEqual(employee.laptop_serial_no, "LAP-100")
        self.assertEqual(employee.birthday.isoformat(), "1996-02-29")
        self.assertTrue(employee.check_password(payload["password"]))
        for label in (b"Address", b"Emergency Contact No", b"Blood Group", b"Laptop serial no", b"Employee ID", b"Birthday"):
            self.assertIn(label, response.data)
        class DetailParser(HTMLParser):
            def __init__(self):
                super().__init__()
                self.fields = []

            def handle_starttag(self, tag, attrs):
                attributes = dict(attrs)
                if tag in ("input", "select", "textarea") and attributes.get("name") in (
                    "address", "emergency_contact_no", "blood_group",
                    "laptop_serial_no", "employee_code", "birthday",
                ):
                    self.fields.append(attributes)

        parser = DetailParser()
        parser.feed(response.data.decode())
        self.assertGreaterEqual(len(parser.fields), 12)
        self.assertTrue(all("required" not in field for field in parser.fields))
        self.assertIn(b'value="1996-02-29"', response.data)

        updated = {**payload, "address": "Updated address", "blood_group": "AB-", "employee_code": "ENG-102", "birthday": "1995-03-15"}
        response = self.client.post(f"/admin/employees/{employee.id}/edit", data=updated, follow_redirects=True)
        self.assertIn(b"Employee details updated.", response.data)
        db.session.refresh(employee)
        self.assertEqual(employee.address, "Updated address")
        self.assertEqual(employee.blood_group, "AB-")
        self.assertEqual(employee.employee_code, "ENG-102")
        self.assertEqual(employee.birthday.isoformat(), "1995-03-15")
        legacy_payload = {key: payload[key] for key in ("name", "email", "mobile", "designation")}
        self.client.post(f"/admin/employees/{employee.id}/edit", data=legacy_payload)
        db.session.refresh(employee)
        self.assertEqual(employee.address, "Updated address")
        self.assertEqual(employee.employee_code, "ENG-102")
        self.assertEqual(employee.birthday.isoformat(), "1995-03-15")
        cleared = {**legacy_payload, **{key: "" for key in (
            "address", "emergency_contact_no", "blood_group", "laptop_serial_no", "employee_code", "birthday"
        )}}
        self.client.post(f"/admin/employees/{employee.id}/edit", data=cleared)
        db.session.refresh(employee)
        self.assertIsNone(employee.employee_code)
        self.assertEqual(employee.address, "")
        self.assertEqual(employee.blood_group, "")
        self.assertIsNone(employee.birthday)

    def test_employee_details_validation_uniqueness_and_privacy(self):
        self._login_admin()
        self.employee.employee_code = "ENG-001"
        self.employee.address = "Private employee address"
        self.employee.emergency_contact_no = "Private emergency number"
        self.employee.blood_group = "AB-"
        self.employee.laptop_serial_no = "PRIVATE-LAPTOP"
        db.session.commit()
        payload = {
            "name": "New Engineer", "email": "new@example.com",
            "mobile": "123", "designation": "Engineer", "password": "password-12345",
        }
        for updates, message in (
            ({"employee_code": " eng-001 "}, b"already uses that Employee ID"),
            ({"blood_group": "X+"}, b"valid blood group"),
            ({"address": "a" * 1001}, b"Address must be at most"),
            ({"emergency_contact_no": "1" * 41}, b"Emergency contact number must be at most"),
            ({"employee_code": "e" * 81}, b"Employee ID must be at most"),
            ({"laptop_serial_no": "s" * 121}, b"Laptop serial number must be at most"),
            ({"birthday": "1995-02-29"}, b"Enter a valid birthday"),
            ({"birthday": "not-a-date"}, b"Enter a valid birthday"),
            ({"birthday": "19960229"}, b"Enter a valid birthday"),
            ({"birthday": "9999-12-31"}, b"Birthday cannot be in the future"),
        ):
            with self.subTest(updates=updates):
                response = self.client.post("/admin/employees", data={**payload, **updates})
                self.assertIn(message, response.data)
                self.assertIsNone(User.query.filter_by(email=payload["email"]).first())
        self.client.post("/admin/employees", data=payload)
        new_employee = User.query.filter_by(email=payload["email"]).one()
        self.assertIsNone(new_employee.birthday)
        response = self.client.post(
            f"/admin/employees/{new_employee.id}/edit",
            data={**payload, "birthday": "invalid", "address": "Should not save"},
            follow_redirects=True,
        )
        self.assertIn(b"Enter a valid birthday", response.data)
        db.session.refresh(new_employee)
        self.assertIsNone(new_employee.birthday)
        self.assertEqual(new_employee.address, "")
        response = self.client.post(
            f"/admin/employees/{new_employee.id}/edit",
            data={**payload, "employee_code": "eng-001", "address": "Should not save"},
            follow_redirects=True,
        )
        self.assertIn(b"already uses that Employee ID", response.data)
        db.session.refresh(new_employee)
        self.assertIsNone(new_employee.employee_code)
        self.assertEqual(new_employee.address, "")
        self.assertIsNone(self.admin.employee_code)
        with self.assertRaises(IntegrityError):
            with db.session.begin_nested():
                new_employee.employee_code = "ENG-001"
                db.session.flush()
        self.client.post("/logout")
        self._login_employee()
        self.assertEqual(self.client.get("/admin/employees").status_code, 403)
        self.assertEqual(self.client.post(
            f"/admin/employees/{new_employee.id}/edit", data=payload
        ).status_code, 403)
        for path in ("/employee", "/api/admin/live-status", "/reports"):
            response = self.client.get(path)
            for private in (b"Private employee address", b"Private emergency number", b"PRIVATE-LAPTOP"):
                self.assertNotIn(private, response.data)

    def test_date_selectors_default_to_current_ist_day_for_both_roles(self):
        self.app.config["SESSION_TIMEOUT_MINUTES"] = 100000
        class DateParser(HTMLParser):
            def __init__(self):
                super().__init__()
                self.fields = {}

            def handle_starttag(self, tag, attrs):
                attributes = dict(attrs)
                if tag == "input" and attributes.get("type") == "date":
                    self.fields[attributes["id"]] = attributes.get("value", "")

        pages = (
            (self._login_admin, {
                "/admin/employees": ["new-birthday", f"edit-{self.employee.id}-birthday"],
                "/admin/leaves/new": ["admin-leave-start", "admin-leave-end"],
                "/admin/leaves": ["filter-from", "filter-to"],
                "/admin/invoices": ["po-start", "po-end", "invoice-po-date", "invoice-start", "invoice-end"],
                "/reports": ["activity-start", "activity-end", "start-date", "end-date"],
            }),
            (self._login_employee, {
                "/employee/leaves": ["leave-start", "leave-end", "filter-from", "filter-to"],
                "/employee/submitted-reports/new": ["report-start", "report-end"],
                "/reports": ["activity-start", "activity-end"],
            }),
        )
        for login, paths in pages:
            self.client.post("/logout")
            login()
            for hour, expected in ((18, "2026-10-29"), (20, "2026-10-30")):
                now = datetime(2026, 10, 29, hour, tzinfo=timezone.utc)
                with patch("geosite.datetime", wraps=datetime) as clock, patch(
                    "geosite.routes.utcnow", return_value=now
                ):
                    clock.now.return_value = now
                    for path, identifiers in paths.items():
                        with self.subTest(path=path, expected=expected):
                            response = self.client.get(path)
                            self.assertEqual(response.status_code, 200)
                            parser = DateParser()
                            parser.feed(response.data.decode())
                            for identifier in identifiers:
                                self.assertEqual(parser.fields[identifier], expected)

    def test_date_defaults_preserve_saved_entered_and_cleared_dates(self):
        self._login_admin()
        self.employee.birthday = date(1996, 2, 29)
        leave = LeaveRequest(
            employee_id=self.employee.id, start_date=date(2026, 9, 1),
            end_date=date(2026, 9, 2), period="full", reason="Historical leave",
        )
        db.session.add(leave)
        db.session.commit()
        response = self.client.get("/admin/employees")
        self.assertIn(b'value="1996-02-29"', response.data)
        response = self.client.get(f"/admin/leaves/{leave.id}/edit")
        self.assertIn(b'value="2026-09-01"', response.data)
        self.assertIn(b'value="2026-09-02"', response.data)
        response = self.client.get("/admin/leaves")
        self.assertNotIn(b"Historical leave", response.data)
        response = self.client.get("/admin/leaves?from_date=2026-09-01&to_date=2026-09-02")
        self.assertIn(b'value="2026-09-01"', response.data)
        self.assertIn(b"Historical leave", response.data)
        response = self.client.get("/admin/leaves?from_date=&to_date=")
        self.assertIn(b'name="from_date" type="date" value=""', response.data)
        self.assertIn(b"Historical leave", response.data)
        response = self.client.post("/admin/leaves/new", data={
            "employee_id": str(self.employee.id), "start_date": "2026-09-03",
            "end_date": "", "period": "full", "reason": "Incomplete dates",
        })
        self.assertEqual(response.status_code, 400)
        self.assertIn(b'value="2026-09-03"', response.data)
        self.assertIn(b'name="end_date" value=""', response.data)
        response = self.client.post("/admin/employees", data={
            "name": "", "birthday": "",
        })
        self.assertIn(b'max="' + india_date(datetime.now(timezone.utc)).isoformat().encode() + b'" value=""', response.data)

    def test_site_navigation_and_add_form_use_site_label(self):
        self._login_admin()
        page = self.client.get("/admin/sites")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"<h1>Site</h1>", page.data)
        self.assertIn(b"Add Site", page.data)
        self.assertNotIn(b"Add Project", page.data)
        self.assertRegex(page.data.decode(), r'(?s)href="/admin/sites"[^>]*>.*?<span>Site</span>')
        self.assertIn(b"Project name", page.data)
        self.assertIn(b'href="/admin/projects"', page.data)

    def test_employee_workspace_uses_admin_assigned_designation(self):
        self._login_admin()
        response = self.client.post(f"/admin/employees/{self.employee.id}/edit", data={
            "name": self.employee.name, "email": self.employee.email,
            "mobile": self.employee.mobile, "designation": "Senior Engineer & Team Lead",
        }, follow_redirects=True)
        self.assertIn(b"Employee details updated.", response.data)
        self.client.post("/logout")
        self._login_employee()
        for path in ("/employee", "/reports", "/employee/leaves"):
            response = self.client.get(path)
            self.assertIn(b"<span>Senior Engineer &amp; Team Lead</span>", response.data)
            self.assertNotIn(b"<span>Field engineer</span>", response.data)
        response = self.client.get("/employee")
        self.assertIn(b"Senior Engineer &amp; Team Lead workspace", response.data)
        self.assertNotIn(b"Field engineer workspace", response.data)
        visit = Attendance(
            employee_id=self.employee.id, site_id=self.site.id, status="onsite",
            checkin_time=datetime.now(timezone.utc), checkin_lat=0, checkin_lon=0,
            checkin_selfie="test.png",
        )
        db.session.add(visit)
        db.session.commit()
        response = self.client.get("/employee")
        self.assertIn(b"Senior Engineer &amp; Team Lead workspace", response.data)
        self.assertIn(b"Onsite now", response.data)

    def test_workspace_navigation_is_role_scoped_and_marks_current_page(self):
        guest = self.client.get("/login")
        self.assertNotIn(b'id="workspaceNav"', guest.data)
        self.assertIn(b'href="#main-content"', guest.data)
        self._login_employee()
        for path, current in (
            ("/employee", "My dashboard"),
            ("/employee/leaves", "My leave"),
            ("/reports", "My reports"),
            ("/submitted-reports", "Submitted reports"),
            ("/employee/submitted-reports/new", "Submitted reports"),
        ):
            with self.subTest(path=path):
                page = self.client.get(path)
                self.assertEqual(page.status_code, 200)
                self.assertIn(b'aria-label="Quick navigation"', page.data)
                self.assertIn(b'aria-controls="workspaceNav"', page.data)
                self.assertNotIn(b'href="/admin/employees"', page.data)
                self.assertRegex(
                    page.data.decode(),
                    rf'(?s)class="workspace-link active"[^>]+aria-current="page".*?<span>{current}</span>',
                )
        self.client.post("/logout")
        self.client.post("/login", data={
            "email": self.admin.email, "password": "admin-password-123",
        })
        page = self.client.get("/admin")
        self.assertIn(b'aria-label="Main navigation"', page.data)
        self.assertIn(b'class="offcanvas offcanvas-start workspace-sidebar"', page.data)
        self.assertIn(b'class="icon-button" type="button" data-bs-toggle="offcanvas"', page.data)
        self.assertIn(b'href="/admin/employees"', page.data)
        self.assertNotIn(b'aria-label="Quick navigation"', page.data)
        self.assertRegex(
            page.data.decode(),
            r'(?s)class="workspace-link active"[^>]+aria-current="page".*?<span>Dashboard</span>',
        )
        live = self.client.get("/api/admin/live-status").get_json()
        self.assertIn("onsite_hours", live["employees"][0])

    def test_attendance_history_shows_each_visit_duration_and_collapsed_locations(self):
        self._login_admin()
        start = datetime(2026, 10, 29, 4, tzinfo=timezone.utc)
        for minutes, status in ((125, "completed"), (66, "report_pending"), (None, "onsite")):
            db.session.add(Attendance(
                employee_id=self.employee.id, site_id=self.site.id,
                checkin_time=start, checkout_time=start + timedelta(minutes=minutes) if minutes else None,
                status=status, checkin_lat=0, checkin_lon=0, checkin_selfie="test.png",
                auto_closed=status == "report_pending",
            ))
        db.session.commit()
        with patch("geosite.routes.utcnow", return_value=start + timedelta(hours=3)):
            page = self.client.get("/admin/attendance")
        self.assertEqual(page.status_code, 200)
        for value in (b"On-site hours", b"2h 05m", b"1h 06m", b"3h 00m", b"In progress", b"Auto-stopped"):
            self.assertIn(value, page.data)
        self.assertEqual(page.data.count(b'<details class="attendance-locations">'), 3)
        self.assertNotIn(b'<details class="attendance-locations" open', page.data)
        self.client.post("/logout")
        self._login_employee()
        self.assertEqual(self.client.get("/admin/attendance").status_code, 403)

    def test_admin_live_status_and_excel(self):
        self.client.post(
            "/login",
            data={"email": self.admin.email, "password": "admin-password-123"},
        )
        self.assertEqual(
            self.client.get("/api/admin/live-status").get_json()["employees"][0]["status"],
            "offline",
        )
        db.session.add(AuditLog(user_id=self.employee.id, action="login"))
        db.session.commit()
        self.assertEqual(
            self.client.get("/api/admin/live-status").get_json()["employees"][0]["status"],
            "assigned",
        )
        result = self.client.get("/admin/reports/export.xlsx")
        self.assertEqual(result.status_code, 302)
        self.assertTrue(result.location.endswith("/reports#activity-export"))
        result = self.client.get("/admin/reports/export.xlsx", query_string={
            "employee_id": self.employee.id, "site_id": self.site.id,
            "start_date": "2026-10-19", "end_date": "2026-10-20",
        })
        self.assertEqual(result.status_code, 200)
        self.assertEqual(load_workbook(io.BytesIO(result.data)).active.title, "SAR-R0")

    def test_password_reset_expires_after_use(self):
        with patch("geosite.routes.send_password_reset_email") as send_email:
            response = self.client.post(
                "/forgot-password",
                data={"email": self.employee.email},
                follow_redirects=True,
            )
        self.assertEqual(response.status_code, 200)
        reset_url = send_email.call_args.args[2]
        token = urlparse(reset_url).path.rsplit("/", 1)[1]
        self.assertEqual(self.client.get(f"/reset-password/{token}").status_code, 200)
        result = self.client.post(
            f"/reset-password/{token}",
            data={
                "password": "employee-new-password-456",
                "confirm_password": "employee-new-password-456",
            },
            follow_redirects=True,
        )
        self.assertEqual(result.status_code, 200)
        self.assertEqual(self.client.get(f"/reset-password/{token}").status_code, 302)
        login = self.client.post(
            "/login",
            data={
                "email": self.employee.email,
                "password": "employee-new-password-456",
            },
        )
        self.assertEqual(login.status_code, 302)


if __name__ == "__main__":
    unittest.main()
