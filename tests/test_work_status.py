import io
import tempfile
import unittest
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

from flask_migrate import downgrade, upgrade
from openpyxl import load_workbook
from sqlalchemy import inspect

from geosite import create_app
from geosite.attendance import daily_cutoff
from geosite.extensions import db
from geosite.models import Attendance, AttendancePolicy, AssignedSite, AuditLog, DailyWorkStatus, LeaveAllowanceAccount, LeaveRequest, MonthlyLeaveEntitlement, SiteAssignment, User
from geosite.timeutils import INDIA_TIME
from geosite.work_status import attendance_summary, month_bounds, monthly_leave_balance


class WorkStatusTests(unittest.TestCase):
    def setUp(self):
        self.uploads = tempfile.TemporaryDirectory()
        self.app = create_app({
            "TESTING": True, "SECRET_KEY": "work-status-tests",
            "SQLALCHEMY_DATABASE_URI": "sqlite://", "UPLOAD_FOLDER": self.uploads.name,
            "WTF_CSRF_ENABLED": False,
        })
        self.context = self.app.app_context()
        self.context.push()
        db.create_all()
        self.employee = self.user("Engineer")
        self.other = self.user("Other")
        self.admin = self.user("Admin", "admin")
        self.site = AssignedSite(site_name="Site", customer_name="Customer", address="Address",
                                 latitude=19, longitude=72, allowed_radius=100)
        db.session.add_all([self.site, AttendancePolicy(id=1, effective_from=date(2026, 10, 26))])
        db.session.flush()
        db.session.add(SiteAssignment(employee=self.employee, site=self.site))
        db.session.commit()
        self.client = self.app.test_client()
        self.now = datetime(2026, 10, 27, 10, tzinfo=timezone.utc)
        self.patches = [
            patch("geosite.routes.utcnow", side_effect=lambda: self.now),
            patch("geosite.work_status.utcnow", side_effect=lambda: self.now),
            patch("geosite.attendance.utcnow", side_effect=lambda: self.now),
        ]
        for item in self.patches:
            item.start()

    def tearDown(self):
        for item in self.patches:
            item.stop()
        db.session.remove()
        db.drop_all()
        db.engine.dispose()
        self.context.pop()
        self.uploads.cleanup()

    def user(self, name, role="employee"):
        user = User(name=name, email=f"{name.lower()}@example.com", mobile=name, role=role,
                    created_date=datetime(2026, 10, 1, tzinfo=timezone.utc))
        user.set_password("work-status-password")
        db.session.add(user)
        db.session.flush()
        return user

    def login(self, user):
        self.client.post("/logout")
        self.client.post("/login", data={"email": user.email, "password": "work-status-password"})
        # Authentication's idle clock is real; attendance uses the fixed policy clock.
        with self.client.session_transaction() as session:
            session["last_activity"] = datetime.now(timezone.utc).isoformat()

    def submit(self, **values):
        return self.client.post("/employee/work-status", data={
            "work_type": "wfh", "reason": "Preparing documentation", **values,
        }, follow_redirects=True)

    def visit(self, when=None, **values):
        visit = Attendance(
            employee_id=self.employee.id, site_id=self.site.id,
            checkin_time=when or self.now, checkin_lat=19, checkin_lon=72,
            checkin_selfie="test.png", status="onsite",
            site_latitude=19, site_longitude=72, site_radius_m=100, **values,
        )
        db.session.add(visit)
        db.session.commit()
        return visit

    def leave(self, start, end=None, status="approved", period="full", employee=None):
        leave = LeaveRequest(employee_id=(employee or self.employee).id, start_date=start,
                             end_date=end or start, status=status, period=period, reason="Leave reason")
        db.session.add(leave)
        db.session.commit()
        return leave

    def summary(self):
        return attendance_summary(self.employee, *month_bounds(date(2026, 10, 27)), now=self.now)

    def test_submit_today_pending_and_admin_approval_removes_automatic_leave(self):
        self.login(self.employee)
        response = self.submit(employee_id=self.other.id, work_date="2026-01-01")
        self.assertIn(b"submitted for administrator approval", response.data)
        item = DailyWorkStatus.query.one()
        self.assertEqual(item.employee_id, self.employee.id)
        self.assertEqual(item.work_date, date(2026, 10, 27))
        self.assertEqual(item.status, "pending")
        self.assertIn(b"WFH", response.data)
        self.assertIsNotNone(AuditLog.query.filter_by(action="work_status_submitted").first())
        self.now = daily_cutoff(self.now)
        self.assertEqual(self.summary()["totals"]["automatic"], 2)
        self.login(self.admin)
        page = self.client.get("/admin/work-status?status=pending")
        self.assertIn(b"Preparing documentation", page.data)
        response = self.client.post(f"/admin/work-status/{item.id}/review",
                                    data={"decision": "approved"}, follow_redirects=True)
        self.assertIn(b"Attendance totals now reflect", response.data)
        db.session.refresh(item)
        self.assertEqual(item.status, "approved")
        self.assertEqual(item.reviewed_by_id, self.admin.id)
        self.assertEqual(self.summary()["totals"]["automatic"], 1)
        self.client.post(f"/admin/work-status/{item.id}/review", data={"decision": "rejected"})
        db.session.refresh(item)
        self.assertEqual(item.status, "approved")
        self.login(self.employee)
        page = self.client.get("/employee")
        self.assertNotIn(b"Automatic leave (days)", page.data)

    def test_admin_dashboard_shows_each_employees_full_daily_status(self):
        today = date(2026, 10, 27)
        self.visit()
        self.leave(today, status="rejected")
        self.other_work_status = DailyWorkStatus(
            employee_id=self.other.id, work_date=today, work_type="wfh",
            reason="Working from home", status="rejected",
            reviewed_by_id=self.admin.id, reviewed_at=self.now,
        )
        training = self.user("Training")
        travelling = self.user("Travelling")
        no_status = self.user("No Status")
        db.session.add_all([
            self.other_work_status,
            DailyWorkStatus(
                employee_id=training.id, work_date=today, work_type="training",
                reason="Safety session", status="approved",
            ),
            DailyWorkStatus(
                employee_id=travelling.id, work_date=today, work_type="travelling",
                reason="Customer visit", status="pending",
            ),
        ])
        db.session.commit()

        self.login(self.admin)
        page = self.client.get("/admin")
        for label in (
            b"Today's status", b"On Site", b"Leave - Full day", b"WFH",
            b"Training", b"Travelling", b"No status submitted",
        ):
            self.assertIn(label, page.data)
        data = self.client.get("/api/admin/live-status").json
        employees = {row["name"]: row for row in data["employees"]}
        self.assertEqual(
            [(item["label"], item["status"]) for item in employees["Engineer"]["daily_statuses"]],
            [("On Site", "onsite"), ("Leave - Full day", "rejected")],
        )
        self.assertTrue(any(
            item["detail"].startswith("Working from home · Reviewed by Admin at ")
            for item in employees["Other"]["daily_statuses"]
        ))
        self.assertEqual(employees["Training"]["daily_statuses"][0]["status"], "approved")
        self.assertEqual(employees["Travelling"]["daily_statuses"][0]["status"], "pending")
        self.assertEqual(
            employees[no_status.name]["daily_statuses"][0]["label"],
            "No status submitted",
        )

    def test_types_reasons_duplicates_and_access_controls(self):
        self.assertEqual(self.submit().status_code, 200)  # Redirected to login.
        self.assertEqual(DailyWorkStatus.query.count(), 0)
        self.login(self.employee)
        for values in ({"work_type": "invalid"}, {"reason": ""}, {"reason": " "}, {"reason": "x" * 2001}):
            self.assertIn(b"Select WFH", self.submit(**values).data)
        self.assertEqual(DailyWorkStatus.query.count(), 0)
        self.submit(work_type="travelling")
        self.assertEqual(DailyWorkStatus.query.one().type_label, "Travelling")
        self.assertIn(b"already been submitted", self.submit(work_type="training").data)
        self.assertEqual(DailyWorkStatus.query.count(), 1)
        self.assertEqual(self.client.get("/admin/work-status").status_code, 403)
        item = DailyWorkStatus.query.one()
        self.assertEqual(self.client.post(f"/admin/work-status/{item.id}/review",
                                         data={"decision": "approved"}).status_code, 403)
        self.login(self.other)
        self.assertNotIn(b"Preparing documentation", self.client.get("/employee").data)
        self.login(self.admin)
        self.assertEqual(self.client.post("/employee/work-status").status_code, 403)
        self.assertEqual(self.client.get("/admin/work-status?status=bad").status_code, 400)
        self.assertEqual(self.client.post(f"/admin/work-status/{item.id}/review",
                                         data={"decision": "bad"}).status_code, 400)
        self.app.config["WTF_CSRF_ENABLED"] = True
        self.assertEqual(self.client.post(f"/admin/work-status/{item.id}/review",
                                         data={"decision": "approved"}).status_code, 400)

    def test_ist_cutoff_monday_to_saturday_policy_and_no_historical_backfill(self):
        # Monday is the policy start; Sunday and earlier dates must never accrue.
        self.now = datetime(2026, 10, 26, 18, 28, 59, tzinfo=timezone.utc)
        self.assertEqual(self.summary()["totals"]["automatic"], 0)
        self.now += timedelta(seconds=1)
        self.assertEqual(self.summary()["totals"]["automatic"], 1)
        self.now = datetime(2026, 11, 1, 18, 29, tzinfo=timezone.utc)  # Sunday.
        self.assertEqual(self.summary()["totals"]["automatic"], 6)
        db.session.get(AttendancePolicy, 1).effective_from = date(2026, 10, 30)
        db.session.commit()
        self.assertEqual(self.summary()["totals"]["automatic"], 2)
        self.employee.created_date = datetime(2026, 10, 31, 0, tzinfo=timezone.utc)
        db.session.commit()
        self.assertEqual(self.summary()["totals"]["automatic"], 1)

    def test_site_visit_even_stopped_without_report_prevents_automatic_leave(self):
        self.visit(when=datetime(2026, 10, 26, 20, tzinfo=timezone.utc))
        self.now = daily_cutoff(self.now)
        # 26 Oct 20:00 UTC is 27 Oct in India.
        self.assertEqual(self.summary()["totals"]["automatic"], 1)
        self.now -= timedelta(seconds=1)
        self.login(self.employee)
        self.assertIn(b"already recorded today", self.submit().data)

    def test_approved_half_leave_is_not_double_counted_and_rejected_work_is_leave(self):
        self.leave(date(2026, 10, 26), period="morning")
        self.now = daily_cutoff(self.now)
        summary = self.summary()
        self.assertEqual(summary["totals"]["approved"], 0.5)
        self.assertEqual(summary["totals"]["automatic"], 1.5)
        self.assertEqual(summary["total_leave"], 2)
        self.leave(date(2026, 10, 26), period="afternoon")
        self.assertEqual(self.summary()["totals"]["automatic"], 1)
        db.session.add(DailyWorkStatus(employee_id=self.employee.id, work_date=date(2026, 10, 27),
                                       work_type="training", reason="Course", status="rejected"))
        db.session.commit()
        self.assertEqual(self.summary()["totals"]["automatic"], 1)

    def test_monthly_totals_clip_cross_month_requests_and_list_all_statuses(self):
        self.leave(date(2026, 9, 29), date(2026, 10, 2))
        self.leave(date(2026, 10, 30), date(2026, 11, 3), status="pending")
        self.leave(date(2026, 10, 3), status="rejected", period="morning")
        self.leave(date(2026, 10, 4), employee=self.other)
        summary = self.summary()
        self.assertEqual(summary["totals"]["approved"], 2)
        self.assertEqual(summary["totals"]["pending"], 2)
        self.assertEqual(summary["totals"]["rejected"], 0.5)
        self.login(self.employee)
        page = self.client.get("/employee")
        self.assertIn(b"My leave - October 2026", page.data)
        for status in ("approved", "pending", "rejected"):
            self.assertIn(f'id="monthly-leave-{status}"'.encode(), page.data)
        self.assertNotIn(b"Automatic leave (days)", page.data)
        self.assertIn(b"Start Site", page.data)
        self.assertNotIn(b"Start site visit", page.data)
        self.visit()
        page = self.client.get("/employee")
        self.assertIn(b"Stop Site", page.data)
        self.assertNotIn(b"End site visit", page.data)

    def test_deadline_rejects_submission_at_2359_and_training_is_valid_before(self):
        self.login(self.employee)
        self.now = daily_cutoff(self.now)
        self.assertIn(b"deadline (23:59 IST) has passed", self.submit().data)
        self.assertEqual(DailyWorkStatus.query.count(), 0)
        self.now -= timedelta(seconds=1)
        self.submit(work_type="training")
        self.assertEqual(DailyWorkStatus.query.one().type_label, "Training")

    def test_monitor_token_scoped_survives_logout_and_detects_stop(self):
        visit = self.visit()
        self.login(self.employee)
        response = self.client.get("/api/employee/active-visit")
        data = response.json
        self.assertEqual(data["attendance_id"], visit.id)
        self.assertEqual(data["radius_m"], 100)
        self.assertEqual(data["latitude"], 19)
        self.assertEqual(data["cutoff_time"], daily_cutoff(self.now).isoformat())
        self.assertIn("no-store", response.headers["Cache-Control"])
        self.site.allowed_radius = 999
        db.session.commit()
        self.assertEqual(self.client.get("/api/employee/active-visit").json["radius_m"], 100)
        token = data["monitoring_token"]
        self.client.post("/logout")
        headers = {"Authorization": f"Bearer {token}"}
        response = self.client.get("/api/visit-monitor", headers=headers)
        self.assertTrue(response.json["active"])
        self.assertNotIn("monitoring_token", response.json)
        visit.status = "report_pending"
        db.session.commit()
        self.assertEqual(self.client.get("/api/visit-monitor", headers=headers).json, {"active": False})

    def test_automatic_leave_flows_to_leave_totals_and_excel_without_double_counting(self):
        self.now = daily_cutoff(self.now)
        self.leave(date(2026, 10, 26), period="morning")
        self.login(self.employee)
        page = self.client.get("/employee/leaves?from_date=2026-10-26&to_date=2026-10-27")
        self.assertIn(b'id="leave-total">2', page.data)
        self.assertIn(b"Automatic leave", page.data)
        params = {"site_id": self.site.id, "start_date": "2026-10-26", "end_date": "2026-10-27"}
        response = self.client.get("/reports/activity.xlsx", query_string=params)
        sheet = load_workbook(io.BytesIO(response.data)).active
        details = [row[4].value for row in sheet.iter_rows(min_row=9)]
        self.assertIn("Automatic leave (0.5 day): no site visit or approved work confirmation.", details)
        self.assertIn("Automatic leave (1 day): no site visit or approved work confirmation.", details)
        self.login(self.admin)
        page = self.client.get(f"/admin/leaves?employee_id={self.employee.id}&from_date=2026-10-26&to_date=2026-10-27")
        self.assertIn(b'id="leave-total">2', page.data)
        params["employee_id"] = self.employee.id
        response = self.client.get("/reports/activity.xlsx", query_string=params)
        self.assertEqual(list(load_workbook(io.BytesIO(response.data)).active.values), list(sheet.values))

    def test_monitor_auth_expiry_tampering_ownership_and_disabled_employee(self):
        self.visit()
        self.login(self.other)
        self.assertEqual(self.client.get("/api/employee/active-visit").json, {"active": False})
        self.login(self.admin)
        self.assertEqual(self.client.get("/api/employee/active-visit").status_code, 403)
        self.login(self.employee)
        token = self.client.get("/api/employee/active-visit").json["monitoring_token"]
        self.client.post("/logout")
        for headers in ({}, {"Authorization": "Bearer bad"}, {"Authorization": f"Bearer {token}x"}):
            self.assertEqual(self.client.get("/api/visit-monitor", headers=headers).status_code, 401)
        self.assertEqual(self.client.get(f"/api/visit-monitor?token={token}").status_code, 401)
        self.employee.is_active = False
        db.session.commit()
        headers = {"Authorization": f"Bearer {token}"}
        self.assertEqual(self.client.get("/api/visit-monitor", headers=headers).status_code, 401)
        self.employee.is_active = True
        db.session.commit()
        self.now = daily_cutoff(self.now)
        self.assertEqual(self.client.get("/api/visit-monitor", headers=headers).status_code, 401)


    def test_daily_selector_assigned_site_only_and_original_location_flow(self):
        second = AssignedSite(site_name="Second site", customer_name="Other customer", address="Address",
                              latitude=20, longitude=73, allowed_radius=100)
        db.session.add(second)
        db.session.flush()
        db.session.add(SiteAssignment(employee=self.employee, site=second))
        db.session.commit()
        self.login(self.employee)
        page = self.client.get("/employee")
        for value in ("onsite", "wfh", "leave", "training", "travelling"):
            self.assertIn(f'value="{value}"'.encode(), page.data)
        for site in (self.site, second):
            self.assertIn(f'<option value="{site.id}">'.encode(), page.data)
            response = self.client.post("/employee/work-status", data={"work_type": "onsite", "site_id": site.id})
            self.assertEqual(response.location, f"/visit/{site.id}/checkin-location")
        self.assertEqual(Attendance.query.count(), 0)
        self.assertIn(b"Select one of your assigned sites", self.submit(work_type="onsite", site_id="").data)
        self.login(self.other)
        self.assertEqual(self.client.post("/employee/work-status", data={
            "work_type": "onsite", "site_id": self.site.id,
        }).status_code, 404)

    def test_daily_leave_requires_reason_and_prevents_duplicates(self):
        self.login(self.employee)
        self.assertIn(b"Invalid leave request", self.submit(work_type="leave", reason=" ").data)
        self.assertEqual(LeaveRequest.query.count(), 0)
        self.assertIn(b"submitted for administrator approval", self.submit(work_type="leave", period="morning").data)
        self.assertEqual(LeaveRequest.query.one().days, 0.5)
        self.assertIn(b"overlap", self.submit(work_type="leave", period="full").data)
        self.assertEqual(LeaveRequest.query.count(), 1)
        self.assertEqual(DailyWorkStatus.query.count(), 0)

    def test_full_day_approved_leave_blocks_all_site_start_surfaces(self):
        self.login(self.employee)
        self.assertEqual(self.client.post(f"/api/visit/{self.site.id}/verify-checkin",
                                         json={"latitude": 19, "longitude": 72}).status_code, 200)
        self.leave(date(2026, 10, 27))
        page = self.client.get("/employee")
        self.assertIn(b"On leave - approved full day", page.data)
        self.assertNotIn(b'id="daily-status-form"', page.data)
        self.assertIn(b"No confirmation is needed", self.submit().data)
        for url in (f"/visit/start/{self.site.id}", f"/visit/{self.site.id}/checkin"):
            response = self.client.post(url, data={}, follow_redirects=True)
            self.assertIn(b"Site starts are blocked", response.data)
        self.assertEqual(Attendance.query.count(), 0)
        self.assertEqual(self.client.get(f"/visit/{self.site.id}/checkin-location").status_code, 302)
        self.assertEqual(self.client.get(f"/visit/{self.site.id}/camera").status_code, 403)
        response = self.client.post(f"/api/visit/{self.site.id}/verify-checkin",
                                    json={"latitude": 19, "longitude": 72})
        self.assertEqual(response.status_code, 409)
        with self.client.session_transaction() as session:
            self.assertNotIn("visit_verification", session)

    def test_half_day_allows_site_start_but_complementary_halves_block(self):
        self.leave(date(2026, 10, 27), period="morning")
        self.login(self.employee)
        page = self.client.get("/employee")
        self.assertIn(b"Approved half-day leave", page.data)
        self.assertIn(b'id="daily-status-form"', page.data)
        self.assertEqual(self.client.post(f"/visit/start/{self.site.id}").location,
                         f"/visit/{self.site.id}/checkin-location")
        self.assertEqual(self.client.post(f"/api/visit/{self.site.id}/verify-checkin",
                                         json={"latitude": 19, "longitude": 72}).status_code, 200)
        self.leave(date(2026, 10, 27), period="afternoon")
        self.assertIn(b"Site starts are blocked", self.client.post(
            f"/visit/{self.site.id}/checkin", follow_redirects=True).data)
        self.assertEqual(Attendance.query.count(), 0)

    def test_admin_add_edit_past_leave_separate_approval_and_audit(self):
        self.login(self.admin)
        self.assertEqual(self.client.get("/admin/leaves/new").status_code, 200)
        values = {"employee_id": str(self.employee.id), "start_date": "2026-10-25",
                  "end_date": "2026-10-25", "period": "full", "reason": "Admin entered leave"}
        response = self.client.post("/admin/leaves/new", data=values, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        leave = LeaveRequest.query.one()
        self.assertEqual(leave.status, "pending")
        self.client.post(f"/admin/leaves/{leave.id}/review", data={"decision": "approved"})
        db.session.refresh(leave)
        self.assertEqual(leave.status, "approved")
        self.assertEqual(self.client.get(f"/admin/leaves/{leave.id}/edit").status_code, 200)
        self.client.post(f"/admin/leaves/{leave.id}/edit", data=values)
        db.session.refresh(leave)
        self.assertEqual(leave.status, "approved")
        self.client.post(f"/admin/leaves/{leave.id}/edit", data={**values, "period": "morning",
                                                               "reason": "Corrected half day"})
        db.session.refresh(leave)
        self.assertEqual(leave.status, "pending")
        self.assertEqual(leave.days, 0.5)
        self.assertIsNone(leave.reviewed_by_id)
        self.assertIsNone(leave.reviewed_at)
        audit = AuditLog.query.filter_by(action="leave_details_saved").order_by(AuditLog.id.desc()).first()
        self.assertIn("Admin entered leave", audit.details)
        self.assertIn("Corrected half day", audit.details)
        self.assertEqual(self.client.post(f"/admin/leaves/{leave.id}/edit",
                                         data={**values, "employee_id": self.other.id}).status_code, 400)
        self.assertEqual(leave.employee_id, self.employee.id)
        self.assertEqual(self.client.post("/admin/leaves/new", data=values).status_code, 400)
        self.assertEqual(self.client.post("/admin/leaves/new", data={**values, "period": "afternoon"}).status_code, 302)

    def test_admin_approval_cannot_make_active_site_full_day_leave(self):
        self.visit()
        morning = self.leave(date(2026, 10, 27), status="pending", period="morning")
        afternoon = self.leave(date(2026, 10, 27), status="pending", period="afternoon")
        self.login(self.admin)
        self.client.post(f"/admin/leaves/{morning.id}/review", data={"decision": "approved"})
        db.session.refresh(morning)
        self.assertEqual(morning.status, "approved")
        response = self.client.post(f"/admin/leaves/{afternoon.id}/review",
                                    data={"decision": "approved"}, follow_redirects=True)
        self.assertIn(b"Stop the employee", response.data)
        db.session.refresh(afternoon)
        self.assertEqual(afternoon.status, "pending")
        self.login(self.employee)
        self.assertIn(b"Stop Site", self.client.get("/employee").data)

    def test_entitlement_upsert_validation_privacy_and_balances(self):
        month = date(2026, 10, 1)
        self.leave(date(2026, 10, 31))
        balance = monthly_leave_balance(self.employee, month, now=self.now)
        self.assertEqual(balance["entitlement"], 2)
        self.assertEqual(balance["remaining"], 1)
        self.assertEqual(balance["used"], 1)  # Monday automatic leave; future approved day is reserved.
        self.assertEqual(balance["reserved"], 1)
        self.login(self.admin)
        values = {"employee_id": self.employee.id, "month": "2026-10", "days": "0.5"}
        self.assertIn(b"Monthly leave rate saved", self.client.post(
            "/admin/leaves/entitlement", data=values, follow_redirects=True).data)
        balance = monthly_leave_balance(self.employee, month, now=self.now)
        self.assertEqual(balance["remaining"], -0.5)
        for invalid in ("NaN", "Infinity", "-1", "31.5", "1.25", "bad", ""):
            self.assertIn(b"Invalid monthly entitlement", self.client.post(
                "/admin/leaves/entitlement", data={**values, "days": invalid}, follow_redirects=True).data)
        for invalid in ("2026-13", "bad", ""):
            self.assertIn(b"Invalid monthly entitlement", self.client.post(
                "/admin/leaves/entitlement", data={**values, "month": invalid}, follow_redirects=True).data)
        self.client.post("/admin/leaves/entitlement", data={**values, "days": "0"})
        self.assertEqual(MonthlyLeaveEntitlement.query.count(), 1)
        self.assertEqual(float(MonthlyLeaveEntitlement.query.one().days), 0)
        self.assertEqual(monthly_leave_balance(self.employee, date(2026, 11, 1), now=self.now)["used"], 1)
        self.assertEqual(monthly_leave_balance(self.employee, date(2026, 11, 1), now=self.now)["entitlement"], 0)
        self.assertEqual(self.client.get("/admin/leaves?leave_month=bad").status_code, 400)
        self.login(self.employee)
        self.assertIn(b'id="monthly-entitlement">0 days', self.client.get("/employee").data)
        self.assertNotIn(self.other.email.encode(), self.client.get("/employee/leaves").data)
        for url in ("/admin/leaves/entitlement", "/admin/leaves/new"):
            self.assertEqual(self.client.post(url, data=values).status_code, 403)
        self.login(self.admin)
        self.app.config["WTF_CSRF_ENABLED"] = True
        self.assertEqual(self.client.post("/admin/leaves/entitlement", data=values).status_code, 400)
        self.assertEqual(self.client.post("/admin/leaves/new", data=values).status_code, 400)

    def test_admin_corrects_work_confirmation_and_employee_cannot_resubmit(self):
        self.login(self.employee)
        self.submit()
        item = DailyWorkStatus.query.one()
        self.assertEqual(self.client.post(f"/admin/work-status/{item.id}/edit", data={}).status_code, 403)
        self.login(self.admin)
        self.client.post(f"/admin/work-status/{item.id}/review", data={"decision": "approved"})
        self.assertEqual(self.client.get(f"/admin/work-status/{item.id}/edit").status_code, 200)
        self.client.post(f"/admin/work-status/{item.id}/edit",
                         data={"work_type": item.work_type, "reason": item.reason})
        db.session.refresh(item)
        self.assertEqual(item.status, "approved")
        self.assertEqual(self.client.post(f"/admin/work-status/{item.id}/edit",
                                         data={"work_type": "training", "reason": ""}).status_code, 400)
        self.client.post(f"/admin/work-status/{item.id}/edit",
                         data={"work_type": "training", "reason": "Corrected by administrator"})
        db.session.refresh(item)
        self.assertEqual(item.work_type, "training")
        self.assertEqual(item.status, "pending")
        self.assertIsNone(item.reviewed_by_id)
        self.assertIsNotNone(AuditLog.query.filter_by(action="work_status_corrected").first())
        self.login(self.employee)
        self.assertIn(b"already been submitted", self.submit().data)
        self.assertEqual(DailyWorkStatus.query.count(), 1)

    def test_two_days_from_joining_month_and_unlimited_year_carryforward(self):
        db.session.delete(db.session.get(AttendancePolicy, 1))
        self.employee.created_date = datetime(2026, 1, 21, 20, tzinfo=timezone.utc)
        db.session.commit()
        for day, expected in ((date(2026, 1, 22), 2), (date(2026, 2, 1), 4),
                              (date(2026, 12, 31), 24), (date(2027, 1, 1), 26)):
            now = datetime.combine(day, datetime.min.time(), tzinfo=INDIA_TIME)
            balance = monthly_leave_balance(self.employee, day, now=now)
            self.assertEqual(balance["entitlement"], expected)
            self.assertEqual(balance["remaining"], expected)
            self.assertEqual(balance["used"], 0)
        self.assertEqual(balance["history"][-1]["credit"], 2)
        self.assertEqual(balance["history"][-2]["remaining"], 24)
        # Future months do not credit ahead of time; months before joining show zero.
        self.assertEqual(monthly_leave_balance(self.employee, date(2028, 1, 1), now=now)["entitlement"], 26)
        self.assertEqual(monthly_leave_balance(self.employee, date(2025, 12, 1), now=now)["entitlement"], 0)
        before_joining = datetime(2026, 1, 21, 19, tzinfo=timezone.utc)  # Joining date is Jan 22 IST.
        self.assertEqual(monthly_leave_balance(self.employee, date(2026, 1, 1),
                                              now=before_joining)["entitlement"], 2)
        before_joining = datetime(2026, 1, 21, 18, tzinfo=timezone.utc)
        self.assertEqual(monthly_leave_balance(self.employee, date(2026, 1, 1),
                                              now=before_joining)["entitlement"], 0)

    def test_opening_balance_rate_changes_and_historical_credits(self):
        db.session.delete(db.session.get(AttendancePolicy, 1))
        db.session.commit()
        self.login(self.admin)
        response = self.client.post("/admin/leaves/account", data={
            "employee_id": self.employee.id, "joining_date": "2026-01-15", "opening_days": "3.5",
        }, follow_redirects=True)
        self.assertIn(b"Joining date and opening leave balance saved", response.data)
        self.client.post("/admin/leaves/entitlement", data={
            "employee_id": self.employee.id, "month": "2026-03", "days": "3",
        })
        self.client.post("/admin/leaves/entitlement", data={
            "employee_id": self.employee.id, "month": "2026-06", "days": "1.5",
        })
        balance = monthly_leave_balance(self.employee, date(2026, 10, 1), now=self.now)
        self.assertEqual(balance["entitlement"], 3.5 + 2 * 2 + 3 * 3 + 5 * 1.5)
        self.assertEqual(balance["history"][0]["credit"], 2)
        self.assertEqual(balance["history"][1]["remaining"], 7.5)
        self.assertEqual(balance["history"][2]["credit"], 3)
        self.assertEqual(balance["monthly_rate"], 1.5)
        self.assertEqual(monthly_leave_balance(self.employee, date(2026, 2, 1), now=self.now)["remaining"], 7.5)
        self.assertEqual(MonthlyLeaveEntitlement.query.count(), 2)
        self.assertIsNotNone(AuditLog.query.filter_by(action="leave_account_updated").first())
        response = self.client.post("/admin/leaves/account", data={
            "employee_id": self.employee.id, "joining_date": "2026-04-01", "opening_days": "3.5",
        }, follow_redirects=True)
        self.assertIn(b"cannot exclude an existing monthly rate change", response.data)
        self.assertEqual(db.session.get(LeaveAllowanceAccount, self.employee.id).joining_date, date(2026, 1, 15))
        self.assertIn(b"Opening balance: 3.5", self.client.get("/admin/leaves?employee_id=1").data)
        self.login(self.employee)
        page = self.client.get("/employee/leaves?employee_id=" + str(self.other.id))
        self.assertIn(b"Jan 2026", page.data)
        self.assertNotIn(self.other.email.encode(), page.data)
        self.assertNotIn(b"Save rate", page.data)
        self.assertNotIn(b"Save joining / opening balance", page.data)
        self.assertEqual(self.client.post("/admin/leaves/account", data={}).status_code, 403)

    def test_account_validation_csrf_and_no_duplicate_opening_credit(self):
        self.login(self.admin)
        values = {"employee_id": self.employee.id, "joining_date": "2026-10-01", "opening_days": "4"}
        for invalid in ("NaN", "Infinity", "-1", "1.25", "1000000", ""):
            response = self.client.post("/admin/leaves/account", data={**values, "opening_days": invalid},
                                        follow_redirects=True)
            self.assertIn(b"Invalid leave account", response.data)
        for invalid in ("bad", "", "2026-10-28"):
            self.assertIn(b"Invalid leave account", self.client.post(
                "/admin/leaves/account", data={**values, "joining_date": invalid},
                follow_redirects=True).data)
        self.assertEqual(LeaveAllowanceAccount.query.count(), 0)
        for _ in range(2):
            self.client.post("/admin/leaves/account", data=values)
        self.assertEqual(LeaveAllowanceAccount.query.count(), 1)
        balance = monthly_leave_balance(self.employee, self.now.date(), now=self.now)
        self.assertEqual(balance["entitlement"], 6)
        self.client.post("/admin/leaves/entitlement", data={
            "employee_id": self.employee.id, "month": "2026-10", "days": "2",
        })
        self.assertIn(b"Invalid leave account", self.client.post("/admin/leaves/account", data={
            **values, "joining_date": "2026-10-28",
        }, follow_redirects=True).data)
        self.assertIn(b"before the employee", self.client.post("/admin/leaves/entitlement", data={
            "employee_id": self.employee.id, "month": "2026-09", "days": "2",
        }, follow_redirects=True).data)
        self.app.config["WTF_CSRF_ENABLED"] = True
        self.assertEqual(self.client.post("/admin/leaves/account", data=values).status_code, 400)

    def test_cumulative_approved_absence_half_days_and_future_reservations(self):
        self.employee.created_date = datetime(2026, 9, 15, tzinfo=timezone.utc)
        db.session.commit()
        self.leave(date(2026, 9, 30))
        self.leave(date(2026, 10, 26), period="morning")
        self.leave(date(2026, 10, 28), period="afternoon")
        self.leave(date(2026, 10, 29), status="pending")
        self.leave(date(2026, 10, 30), status="rejected")
        balance = monthly_leave_balance(self.employee, date(2026, 10, 1), now=self.now)
        self.assertEqual(balance["entitlement"], 4)
        self.assertEqual(balance["used"], 2)  # September 1 + October half approved + half absence.
        self.assertEqual(balance["remaining"], 2)
        self.assertEqual(balance["reserved"], 0.5)
        self.assertEqual(balance["available"], 1.5)
        self.assertEqual(balance["history"][0]["approved"], 1)
        self.assertEqual(balance["history"][1]["automatic"], 0.5)
        self.assertEqual(sum(row["approved"] + row["automatic"] for row in balance["history"]), balance["used"])
        self.now = datetime(2026, 10, 28, 10, tzinfo=timezone.utc)
        balance = monthly_leave_balance(self.employee, date(2026, 10, 1), now=self.now)
        self.assertEqual(balance["used"], 3.5)  # Tuesday absent; Wednesday approved half.
        self.assertEqual(balance["reserved"], 0)
        self.assertEqual(balance["remaining"], 0.5)


class WorkStatusMigrationTests(unittest.TestCase):
    def test_upgrade_policy_activation_and_reversible_schema(self):
        with tempfile.TemporaryDirectory() as directory:
            app = create_app({
                "TESTING": True, "SECRET_KEY": "work-status-migration",
                "SQLALCHEMY_DATABASE_URI": "sqlite:///" + str(Path(directory) / "work.db"),
                "UPLOAD_FOLDER": str(Path(directory) / "uploads"),
            })
            migrations = str(Path(__file__).resolve().parents[1] / "migrations")
            with app.app_context():
                try:
                    upgrade(directory=migrations, revision="e48f9a0b1c27")
                    user = User(name="Old", email="old@example.com", mobile="1", password_hash="test")
                    db.session.add(user)
                    db.session.commit()
                    upgrade(directory=migrations)
                    self.assertEqual(db.session.get(AttendancePolicy, 1).effective_from,
                                     datetime.now(INDIA_TIME).date())
                    self.assertEqual(DailyWorkStatus.query.count(), 0)
                    self.assertIn("monthly_leave_entitlement", inspect(db.engine).get_table_names())
                    self.assertIn("leave_allowance_account", inspect(db.engine).get_table_names())
                    self.assertEqual(MonthlyLeaveEntitlement.query.count(), 0)
                    self.assertIn("site_radius_m", {
                        column["name"] for column in inspect(db.engine).get_columns("attendance")
                    })
                    user = User.query.one()
                    user.created_date = datetime(2026, 1, 1, tzinfo=timezone.utc)
                    db.session.add(MonthlyLeaveEntitlement(
                        employee_id=user.id, month=date(2026, 2, 1), days=3,
                        updated_by_id=user.id,
                    ))
                    db.session.commit()
                    db.session.remove()
                    downgrade(directory=migrations, revision="06ab1c2d3e49")
                    self.assertNotIn("leave_allowance_account", inspect(db.engine).get_table_names())
                    self.assertEqual(float(MonthlyLeaveEntitlement.query.one().days), 3)
                    db.session.remove()
                    upgrade(directory=migrations)
                    self.assertEqual(LeaveAllowanceAccount.query.count(), 0)
                    balance = monthly_leave_balance(
                        User.query.one(), date(2026, 3, 1),
                        now=datetime(2026, 3, 15, tzinfo=timezone.utc),
                    )
                    self.assertEqual(balance["entitlement"], 8)  # January 2, February/March 3 each.
                    db.session.remove()
                    downgrade(directory=migrations, revision="e48f9a0b1c27")
                    self.assertNotIn("daily_work_status", inspect(db.engine).get_table_names())
                    self.assertNotIn("monthly_leave_entitlement", inspect(db.engine).get_table_names())
                    self.assertNotIn("leave_allowance_account", inspect(db.engine).get_table_names())
                    self.assertEqual(User.query.one().name, "Old")
                    upgrade(directory=migrations)
                    self.assertEqual(DailyWorkStatus.query.count(), 0)
                finally:
                    db.session.remove()
                    db.engine.dispose()


if __name__ == "__main__":
    unittest.main()
